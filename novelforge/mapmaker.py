"""
Fantasy map maker - data model, renderers and export.

The architecture that matters: a map is turned into a flat list of drawing
**primitives** by `build_primitives`. The Tkinter editor draws that list onto a
Canvas; the PNG exporter draws the same list with Pillow; the SVG exporter
writes the same list as XML. One display list, three backends, so what you see
while editing is exactly what you export.

Maps live in `13 Maps` as `<name>.map.json` alongside their exports. They are
deliberately kept out of project.json - a hand-drawn coastline is thousands of
points and would bloat the manifest that gets rewritten on every keystroke.
"""

from __future__ import annotations

import json
import math
import os
import random
import re
import zlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from .atomic import read_json, safe_filename, write_json_atomic
from .model import from_dict, new_id, now_iso, to_dict

Point = Tuple[float, float]

# ==========================================================================
# Styles
# ==========================================================================

# Each style is a full palette.
#
# `paper` is the SEA - the page a map is drawn on reads as ocean once land is
# painted over it - so `terrain.land` must be clearly lighter than `paper` or
# the coastline reads as a bare outline with no land inside it. That contrast
# is the single most important thing separating a map from a diagram.
#
# `halo` is the concentric coastal glow, drawn under the land fill in
# progressively wider, lighter strokes. `terrain` overrides the per-kind
# defaults in TERRAIN so every style stays internally consistent.
STYLES: Dict[str, Dict[str, Any]] = {
    "parchment": {
        "label": "Parchment",
        "paper": "#d9c398",
        "paper_dark": "#bfa374",
        "ink": "#4a3520",
        "ink_light": "#7a5f3e",
        "water": "#a8c2ca",
        "water_deep": "#8fb0ba",
        "halo": ["#c9b48b", "#d2bf99", "#dbcaa8"],
        "grid": "#c4b18c",
        "texture": True,
        "vignette": True,
        "terrain": {
            "land": "#f2e8cb", "water": "#a8c2ca", "forest": "#b3c294",
            "desert": "#eadfae", "swamp": "#adb488", "ice": "#e4ebeb",
        },
    },
    "ink": {
        "label": "Ink on white",
        "paper": "#e9eef1",
        "paper_dark": "#d8e0e4",
        "ink": "#1e1e1e",
        "ink_light": "#5a5a5a",
        "water": "#d3e0e6",
        "water_deep": "#bfd2da",
        "halo": ["#dce5e9", "#e3eaee", "#eaeff2"],
        "grid": "#cbd4d8",
        "texture": False,
        "vignette": False,
        "terrain": {
            "land": "#fdfdfb", "water": "#d3e0e6", "forest": "#dde8d5",
            "desert": "#f5efdc", "swamp": "#dfe2cf", "ice": "#f2f7f8",
        },
    },
    "dark": {
        "label": "Dark atlas",
        "paper": "#161c22",
        "paper_dark": "#101519",
        "ink": "#ddd6c6",
        "ink_light": "#8f887c",
        "water": "#1b2731",
        "water_deep": "#141d25",
        "halo": ["#222c35", "#1d262e", "#192128"],
        "grid": "#2b343c",
        "texture": True,
        "vignette": True,
        "terrain": {
            "land": "#333c44", "water": "#1b2731", "forest": "#2f4038",
            "desert": "#45412f", "swamp": "#333a2e", "ice": "#3f4a52",
        },
    },
    "treasure": {
        "label": "Treasure map",
        "paper": "#c9a86f",
        "paper_dark": "#a88a52",
        "ink": "#5a3a18",
        "ink_light": "#8a6535",
        "water": "#b9a06d",
        "water_deep": "#a48c5c",
        "halo": ["#bd9c63", "#c6a771", "#cfb281"],
        "grid": "#b5975f",
        "texture": True,
        "vignette": True,
        "terrain": {
            "land": "#e8d3a4", "water": "#b9a06d", "forest": "#b0ac78",
            "desert": "#e2cd97", "swamp": "#a5a274", "ice": "#dcdcc9",
        },
    },
}

TERRAIN: Dict[str, Dict[str, Any]] = {
    "land":      {"label": "Land / coast", "fill": "#f2e8cb", "closed": True,
                  "halo": True, "decor": None},
    "water":     {"label": "Sea / lake", "fill": "#a8c2ca", "closed": True,
                  "halo": False, "decor": None},
    "forest":    {"label": "Forest", "fill": "#b3c294", "closed": True,
                  "halo": False, "decor": "trees"},
    "mountains": {"label": "Mountain range", "fill": None, "closed": False,
                  "halo": False, "decor": "peaks"},
    "hills":     {"label": "Hills", "fill": None, "closed": False,
                  "halo": False, "decor": "hills"},
    "desert":    {"label": "Desert", "fill": "#e4d5a8", "closed": True,
                  "halo": False, "decor": "dots"},
    "swamp":     {"label": "Marsh / swamp", "fill": "#b3b585", "closed": True,
                  "halo": False, "decor": "squiggles"},
    "ice":       {"label": "Ice / tundra", "fill": "#dfe7e8", "closed": True,
                  "halo": False, "decor": None},
    "region":    {"label": "Region / border", "fill": None, "closed": True,
                  "halo": False, "decor": None},
    "river":     {"label": "River", "fill": None, "closed": False,
                  "halo": False, "decor": "river"},
    "road":      {"label": "Road", "fill": None, "closed": False,
                  "halo": False, "decor": "road"},
    "wall":      {"label": "Wall", "fill": None, "closed": False,
                  "halo": False, "decor": "wall"},
    "route":     {"label": "Route / journey", "fill": None, "closed": False,
                  "halo": False, "decor": "route"},
}

TERRAIN_ORDER = [
    "land", "water", "forest", "mountains", "hills", "desert", "swamp",
    "ice", "region", "river", "road", "wall", "route",
]

# Draw order, low to high; shapes of equal rank keep the order they were made
# in. Water and land share a rank on purpose: a sea drawn first lies under the
# land drawn after it, and a lake drawn on the land lies on top of it. With
# water always underneath, every lake a writer drew simply vanished.
PAINT_ORDER = {
    "water": 1, "land": 1, "ice": 2, "desert": 3, "swamp": 4, "forest": 5,
    "hills": 6, "mountains": 7, "region": 8, "river": 9, "road": 10,
    "wall": 11, "route": 12,
}

PIN_KINDS: Dict[str, str] = {
    "capital": "Capital city",
    "city": "City",
    "town": "Town",
    "village": "Village",
    "castle": "Castle / keep",
    "tower": "Tower",
    "temple": "Temple",
    "port": "Port / harbour",
    "ruin": "Ruin",
    "cave": "Cave",
    "dungeon": "Dungeon",
    "mine": "Mine",
    "camp": "Camp",
    "bridge": "Bridge",
    "inn": "Inn",
    "battle": "Battle site",
    "danger": "Danger",
    "treasure": "Treasure",
    "landmark": "Landmark",
    "portal": "Portal",
}

MAP_KINDS = {
    "world": "World",
    "continent": "Continent",
    "region": "Region",
    "city": "City / town plan",
    "building": "Building / interior",
    "dungeon": "Dungeon",
    "treasure": "Treasure map",
    "battle": "Battle plan",
}

GRID_KINDS = {"none": "None", "square": "Square", "hex": "Hex"}

#: Who a picture is for. "author" is the working copy (every visible layer,
#: notes and all); "reader" is what leaves the building - a book, an ebook, a
#: Word file for a proofreader - and drops every author-only layer and every
#: note.
EDITIONS = ("author", "reader")


# ==========================================================================
# Model
# ==========================================================================


@dataclass
class Layer:
    name: str = "Base"
    visible: bool = True
    locked: bool = False
    # Secrets: a layer the author keeps for themselves (unfinished places, plot
    # spoilers, GM notes). Everything drawn for the *reader edition* leaves it
    # out - the picture, the SVG, the Word listing, the alt text.
    author_only: bool = False


@dataclass
class Shape:
    id: str = ""
    kind: str = "land"
    points: List[Point] = field(default_factory=list)
    closed: bool = True
    fill: str = ""          # blank means "use the terrain default"
    outline: str = ""
    width: float = 2.0
    label: str = ""
    layer: str = "Base"

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("shp")
        # JSON round-trips tuples into lists; normalise so geometry maths works.
        self.points = [(float(p[0]), float(p[1])) for p in self.points if len(p) >= 2]

    def bounds(self) -> Tuple[float, float, float, float]:
        if not self.points:
            return (0.0, 0.0, 0.0, 0.0)
        xs = [p[0] for p in self.points]
        ys = [p[1] for p in self.points]
        return (min(xs), min(ys), max(xs), max(ys))

    def centroid(self) -> Point:
        if not self.points:
            return (0.0, 0.0)
        return (sum(p[0] for p in self.points) / len(self.points),
                sum(p[1] for p in self.points) / len(self.points))


@dataclass
class Pin:
    id: str = ""
    x: float = 0.0
    y: float = 0.0
    kind: str = "city"
    label: str = ""
    label_side: str = "e"      # e | w | n | s
    entity_id: str = ""        # links to a Location entity in the project
    notes: str = ""
    size: float = 7.0
    layer: str = "Base"
    # The map that shows this place from the inside (a village under a world
    # pin, a floor plan under a town pin). Empty when there is none.
    child_map_id: str = ""

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("pin")


@dataclass
class MapLabel:
    id: str = ""
    x: float = 0.0
    y: float = 0.0
    text: str = ""
    size: int = 16
    color: str = ""
    italic: bool = False
    bold: bool = False
    tracking: float = 0.0      # extra letter spacing, for region names
    layer: str = "Base"

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("lbl")


@dataclass
class GameMap:
    id: str = ""
    name: str = "New Map"
    kind: str = "world"
    style: str = "parchment"
    width: int = 1600
    height: int = 1100
    grid: str = "none"
    grid_size: int = 80
    scale_text: str = ""
    compass: bool = True
    border: bool = True
    title_on_map: bool = True
    # When a place name has nowhere to go, hide it rather than print it on top
    # of another. Turn off to see every name, overlaps and all.
    hide_colliding_labels: bool = True
    auto_place_labels: bool = True
    notes: str = ""
    seed: int = 7
    layers: List[Layer] = field(default_factory=lambda: [Layer()])
    shapes: List[Shape] = field(default_factory=list)
    pins: List[Pin] = field(default_factory=list)
    labels: List[MapLabel] = field(default_factory=list)
    created: str = field(default_factory=now_iso)
    modified: str = field(default_factory=now_iso)
    # How long the scale bar is drawn, in map pixels. 0 means "the default for
    # this map's width" (see `scale_bar`); `apply_scale` sets it so the bar can
    # be a round number of units. The picture, the label placer and every
    # distance NovelForge reports all read it through the one `scale_bar`.
    scale_px: float = 0.0
    # An on-map key: a swatch and a name for every kind of thing drawn.
    legend: bool = False
    # For a map that is a pin's inside (a city under a world pin, a floor plan
    # under a city pin): {"map_id": ..., "pin_id": ...}. Empty for a top map.
    parent: Dict[str, str] = field(default_factory=dict)

    # Not part of the map: a memo of the last label-placement pass, so panning
    # and zooming (which change nothing about where pins and labels sit) do
    # not re-run that collision search on every single redraw. Never
    # serialised - see to_json below.
    _label_cache: Optional[Tuple[tuple, Dict[str, Tuple[str, bool]]]] = field(
        default=None, init=False, repr=False, compare=False)
    # The same idea for the whole display list, and for each shape's share of
    # it: a redraw caused by zooming, panning or selecting something changes
    # nothing about the map, so it should not rebuild thousands of trees.
    _prim_cache: Optional[Tuple[tuple, List[tuple]]] = field(
        default=None, init=False, repr=False, compare=False)
    _shape_cache: Dict[tuple, Tuple[tuple, List[tuple]]] = field(
        default_factory=dict, init=False, repr=False, compare=False)

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("map")
        if not self.layers:
            self.layers = [Layer()]

    # -- serialisation --------------------------------------------------
    def to_json(self) -> Dict[str, Any]:
        return {
            "id": self.id, "name": self.name, "kind": self.kind,
            "style": self.style, "width": self.width, "height": self.height,
            "grid": self.grid, "grid_size": self.grid_size,
            "scale_text": self.scale_text, "compass": self.compass,
            "border": self.border, "title_on_map": self.title_on_map,
            "hide_colliding_labels": self.hide_colliding_labels,
            "auto_place_labels": self.auto_place_labels,
            "notes": self.notes, "seed": self.seed,
            "scale_px": self.scale_px, "legend": self.legend,
            "parent": dict(self.parent or {}),
            "layers": [to_dict(l) for l in self.layers],
            "shapes": [
                {**to_dict(s), "points": [[p[0], p[1]] for p in s.points]}
                for s in self.shapes
            ],
            "pins": [to_dict(p) for p in self.pins],
            "labels": [to_dict(l) for l in self.labels],
            "created": self.created, "modified": self.modified,
        }

    @classmethod
    def from_json(cls, data: Dict[str, Any]) -> "GameMap":
        skip = {"layers", "shapes", "pins", "labels"}
        obj = cls(**{
            k: v for k, v in (data or {}).items()
            if k not in skip and k in {f.name for f in cls.__dataclass_fields__.values()}
        })
        obj.layers = [from_dict(Layer, d) for d in (data.get("layers") or [])] \
            or [Layer()]
        obj.shapes = [from_dict(Shape, d) for d in (data.get("shapes") or [])]
        obj.pins = [from_dict(Pin, d) for d in (data.get("pins") or [])]
        obj.labels = [from_dict(MapLabel, d) for d in (data.get("labels") or [])]
        # Files from before these fields existed simply lack them; a hand-edited
        # one might hold nonsense. Either way the map should still open.
        if not isinstance(obj.parent, dict):
            obj.parent = {}
        try:
            obj.scale_px = max(0.0, float(obj.scale_px or 0.0))
        except (TypeError, ValueError):
            obj.scale_px = 0.0
        obj.legend = bool(obj.legend)
        return obj

    # -- helpers --------------------------------------------------------
    def palette(self) -> Dict[str, Any]:
        return STYLES.get(self.style, STYLES["parchment"])

    def layer_names(self) -> List[str]:
        return [l.name for l in self.layers]

    def visible_layers(self, edition: str = "author") -> set:
        """
        Names of the layers that are drawn.

        `edition="reader"` is the copy that leaves the building - a printed
        book, an ebook, a Word file for a proofreader - so it also drops every
        layer marked author-only. The author's own view keeps them.
        """
        if edition not in EDITIONS:
            # A typo must not quietly fall back to the author's copy.
            raise ValueError(f"edition must be one of {EDITIONS}, not {edition!r}")
        reader = edition == "reader"
        return {l.name for l in self.layers
                if l.visible and not (reader and l.author_only)}

    def layer(self, name: str) -> Optional[Layer]:
        return next((l for l in self.layers if l.name == name), None)

    def shape(self, shape_id: str) -> Optional[Shape]:
        return next((s for s in self.shapes if s.id == shape_id), None)

    def pin(self, pin_id: str) -> Optional[Pin]:
        return next((p for p in self.pins if p.id == pin_id), None)

    def label(self, label_id: str) -> Optional[MapLabel]:
        return next((l for l in self.labels if l.id == label_id), None)

    def touch(self) -> None:
        self.modified = now_iso()

    def is_empty(self) -> bool:
        return not (self.shapes or self.pins or self.labels)


# ==========================================================================
# Primitives - the shared display list
# ==========================================================================

# ("polygon", points, fill, outline, width, dash)
# ("line",    points, color, width, dash)
# ("ellipse", x0, y0, x1, y1, fill, outline, width)
# ("text",    x, y, text, size, color, anchor, italic, bold, tracking[, halo])
#
# `halo` is optional: a colour to outline the letters with, so a name stays
# legible where it crosses a coast, a river or a mountain. Every backend reads
# the list through `text_parts`, so none of them breaks on the short form.


def _rgb(colour: str) -> Tuple[int, int, int]:
    colour = (colour or "#000000").lstrip("#")
    if len(colour) == 3:
        colour = "".join(c * 2 for c in colour)
    if len(colour) != 6:
        return (0, 0, 0)
    try:
        return (int(colour[0:2], 16), int(colour[2:4], 16), int(colour[4:6], 16))
    except ValueError:
        return (0, 0, 0)


def _mix(a: str, b: str, t: float) -> str:
    ra, ga, ba = _rgb(a)
    rb, gb, bb = _rgb(b)
    return "#%02x%02x%02x" % (
        int(ra + (rb - ra) * t), int(ga + (gb - ga) * t), int(ba + (bb - ba) * t)
    )


def _smooth(points: Sequence[Point], iterations: int = 2) -> List[Point]:
    """
    Chaikin corner-cutting, so a hand-drawn coastline is not visibly polygonal.

    Bounded by `iterations` and skipped for very long paths, because each pass
    doubles the point count and a freehand coast can already hold thousands.
    """
    pts = [(float(p[0]), float(p[1])) for p in points]
    if len(pts) < 3 or len(pts) > 600:
        return pts
    for _ in range(max(0, min(3, iterations))):
        out: List[Point] = [pts[0]]
        for a, b in zip(pts, pts[1:]):
            out.append((a[0] * 0.75 + b[0] * 0.25, a[1] * 0.75 + b[1] * 0.25))
            out.append((a[0] * 0.25 + b[0] * 0.75, a[1] * 0.25 + b[1] * 0.75))
        out.append(pts[-1])
        pts = out
        if len(pts) > 2400:
            break
    return pts


def _path_length(points: Sequence[Point]) -> float:
    return sum(
        math.hypot(b[0] - a[0], b[1] - a[1]) for a, b in zip(points, points[1:])
    )


def _walk_path(points: Sequence[Point], spacing: float) -> List[Tuple[Point, float]]:
    """
    Sample points along a polyline at roughly `spacing` apart.

    Returns (point, angle). Used to scatter mountain peaks and trees along a
    range. `spacing` is clamped so a degenerate value cannot loop forever.
    """
    spacing = max(4.0, float(spacing))
    out: List[Tuple[Point, float]] = []
    carry = 0.0
    for a, b in zip(points, points[1:]):
        dx, dy = b[0] - a[0], b[1] - a[1]
        seg = math.hypot(dx, dy)
        if seg < 1e-6:
            continue
        angle = math.atan2(dy, dx)
        position = carry
        while position < seg:
            t = position / seg
            out.append(((a[0] + dx * t, a[1] + dy * t), angle))
            position += spacing
            if len(out) > 4000:      # hard cap: decoration, not geometry
                return out
        carry = position - seg
    return out


def _stable_seed(shape: Shape) -> int:
    """
    A number that is the same on every launch.

    Python salts `hash()` of a string differently in every process, so seeding
    with it reshuffled every forest each time a map was opened - the same map
    looked different on Monday and on Tuesday.
    """
    return zlib.crc32(shape.id.encode("utf-8")) & 0xFFFF


def _inside_test(shape: Shape, x0: float, y0: float, x1: float, y1: float):
    """
    A fast "is this point inside the shape?" test.

    Ray casting is O(points) per question, and a forest asks about thousands of
    spots against a coast of hundreds of points. Drawing the polygon once into
    a small bitmap and looking pixels up is a hundred times quicker, and one
    pixel of error at the edge of a forest is invisible.
    """
    points = shape.points
    if len(points) > 24:
        try:
            from PIL import Image, ImageDraw

            span = max(x1 - x0, y1 - y0, 1.0)
            scale = min(1.0, 360.0 / span)
            mask = Image.new("L", (max(2, int((x1 - x0) * scale) + 2),
                                   max(2, int((y1 - y0) * scale) + 2)), 0)
            ImageDraw.Draw(mask).polygon(
                [((px - x0) * scale, (py - y0) * scale) for px, py in points],
                fill=255)
            pixels = mask.load()
            mw, mh = mask.size

            def inside(px: float, py: float) -> bool:
                ix, iy = int((px - x0) * scale), int((py - y0) * scale)
                return 0 <= ix < mw and 0 <= iy < mh and pixels[ix, iy] > 0

            return inside
        except Exception:
            pass
    return lambda px, py: point_in_polygon((px, py), points)


def _scatter_in_bounds(shape: Shape, spacing: float, seed: int,
                       limit: int = 6000) -> List[Point]:
    """Jittered grid of points inside a shape's bounding box, filtered by the polygon."""
    x0, y0, x1, y1 = shape.bounds()
    if x1 - x0 < 2 or y1 - y0 < 2:
        return []
    rng = random.Random(seed ^ _stable_seed(shape))
    spacing = max(10.0, float(spacing))
    out: List[Point] = []
    rows = int((y1 - y0) / spacing) + 1
    cols = int((x1 - x0) / spacing) + 1
    if rows * cols > limit:           # keep decoration bounded on huge shapes
        spacing = math.sqrt((x1 - x0) * (y1 - y0) / float(limit))
        rows = int((y1 - y0) / spacing) + 1
        cols = int((x1 - x0) / spacing) + 1
    inside = _inside_test(shape, x0, y0, x1, y1)
    for r in range(rows):
        for c in range(cols):
            px = x0 + c * spacing + rng.uniform(-spacing * 0.3, spacing * 0.3)
            py = y0 + r * spacing + rng.uniform(-spacing * 0.3, spacing * 0.3)
            if inside(px, py):
                out.append((px, py))
    return out


def point_in_polygon(pt: Point, polygon: Sequence[Point]) -> bool:
    """Ray casting. Used for hit-testing and for filling terrain decoration."""
    if len(polygon) < 3:
        return False
    x, y = pt
    inside = False
    n = len(polygon)
    j = n - 1
    for i in range(n):
        xi, yi = polygon[i]
        xj, yj = polygon[j]
        if ((yi > y) != (yj > y)) and \
                (x < (xj - xi) * (y - yi) / ((yj - yi) or 1e-12) + xi):
            inside = not inside
        j = i
    return inside


def distance_to_path(pt: Point, points: Sequence[Point]) -> float:
    """Shortest distance from a point to a polyline, for hit-testing lines."""
    if not points:
        return float("inf")
    if len(points) == 1:
        return math.hypot(pt[0] - points[0][0], pt[1] - points[0][1])
    best = float("inf")
    px, py = pt
    for a, b in zip(points, points[1:]):
        ax, ay = a
        bx, by = b
        dx, dy = bx - ax, by - ay
        seg_sq = dx * dx + dy * dy
        if seg_sq < 1e-12:
            best = min(best, math.hypot(px - ax, py - ay))
            continue
        t = max(0.0, min(1.0, ((px - ax) * dx + (py - ay) * dy) / seg_sq))
        best = min(best, math.hypot(px - (ax + t * dx), py - (ay + t * dy)))
    return best


# -- pin glyphs ------------------------------------------------------------


def pin_primitives(pin: Pin, ink: str, paper: str,
                   accent: str = "") -> List[tuple]:
    """
    Build the drawing primitives for one map pin.

    Shared by the editor and both exporters, so a treasure X on screen is the
    same treasure X in the PNG.
    """
    x, y = float(pin.x), float(pin.y)
    s = max(3.0, float(pin.size))
    kind = pin.kind
    out: List[tuple] = []
    accent = accent or ink

    def ring(radius: float, fill: str, width: float = 1.6) -> None:
        out.append(("ellipse", x - radius, y - radius, x + radius, y + radius,
                    fill, ink, width))

    if kind == "capital":
        ring(s * 1.55, paper, 1.8)
        ring(s * 0.75, ink, 1.2)
        for i in range(8):                      # star burst
            a = i * math.pi / 4
            out.append(("line",
                        [(x + math.cos(a) * s * 1.6, y + math.sin(a) * s * 1.6),
                         (x + math.cos(a) * s * 2.15, y + math.sin(a) * s * 2.15)],
                        ink, 1.4, False))
    elif kind == "city":
        ring(s * 1.15, paper, 1.8)
        ring(s * 0.45, ink, 1.0)
    elif kind == "town":
        ring(s * 0.8, paper, 1.6)
        out.append(("ellipse", x - s * 0.22, y - s * 0.22,
                    x + s * 0.22, y + s * 0.22, ink, ink, 1.0))
    elif kind == "village":
        out.append(("polygon",
                    [(x - s * 0.6, y + s * 0.5), (x + s * 0.6, y + s * 0.5),
                     (x + s * 0.6, y - s * 0.15), (x, y - s * 0.75),
                     (x - s * 0.6, y - s * 0.15)],
                    paper, ink, 1.4, False))
    elif kind == "castle":
        base_y = y + s * 0.7
        top_y = y - s * 0.35
        out.append(("polygon",
                    [(x - s, base_y), (x - s, top_y), (x - s * 0.6, top_y),
                     (x - s * 0.6, top_y - s * 0.4), (x - s * 0.2, top_y - s * 0.4),
                     (x - s * 0.2, top_y), (x + s * 0.2, top_y),
                     (x + s * 0.2, top_y - s * 0.4), (x + s * 0.6, top_y - s * 0.4),
                     (x + s * 0.6, top_y), (x + s, top_y), (x + s, base_y)],
                    paper, ink, 1.5, False))
    elif kind == "tower":
        out.append(("polygon",
                    [(x - s * 0.45, y + s * 0.9), (x - s * 0.45, y - s * 0.5),
                     (x, y - s * 1.15), (x + s * 0.45, y - s * 0.5),
                     (x + s * 0.45, y + s * 0.9)],
                    paper, ink, 1.5, False))
    elif kind == "temple":
        out.append(("polygon",
                    [(x - s, y + s * 0.3), (x, y - s * 0.95), (x + s, y + s * 0.3)],
                    paper, ink, 1.5, False))
        out.append(("line", [(x - s, y + s * 0.7), (x + s, y + s * 0.7)],
                    ink, 1.6, False))
    elif kind == "port":
        out.append(("line", [(x, y - s), (x, y + s * 0.9)], ink, 1.8, False))
        out.append(("line", [(x - s * 0.7, y - s * 0.55),
                             (x + s * 0.7, y - s * 0.55)], ink, 1.6, False))
        out.append(("line", [(x - s * 0.85, y + s * 0.25),
                             (x, y + s * 0.95), (x + s * 0.85, y + s * 0.25)],
                    ink, 1.7, False))
    elif kind == "ruin":
        out.append(("line", [(x - s * 0.8, y + s * 0.7),
                             (x - s * 0.8, y - s * 0.5)], ink, 1.8, False))
        out.append(("line", [(x - s * 0.1, y + s * 0.7),
                             (x - s * 0.1, y - s * 0.1)], ink, 1.8, False))
        out.append(("line", [(x + s * 0.7, y + s * 0.7),
                             (x + s * 0.7, y - s * 0.65)], ink, 1.8, False))
        out.append(("line", [(x - s, y + s * 0.8), (x + s, y + s * 0.8)],
                    ink, 1.5, False))
    elif kind in ("cave", "dungeon"):
        out.append(("polygon",
                    [(x - s, y + s * 0.75), (x - s * 0.75, y - s * 0.35),
                     (x, y - s * 0.8), (x + s * 0.75, y - s * 0.35),
                     (x + s, y + s * 0.75)],
                    ink if kind == "dungeon" else paper, ink, 1.5, False))
        if kind == "dungeon":
            out.append(("line", [(x - s * 0.35, y + s * 0.75),
                                 (x - s * 0.35, y + s * 0.1)], paper, 1.4, False))
            out.append(("line", [(x + s * 0.35, y + s * 0.75),
                                 (x + s * 0.35, y + s * 0.1)], paper, 1.4, False))
    elif kind == "mine":
        out.append(("polygon",
                    [(x - s, y + s * 0.7), (x, y - s * 0.8), (x + s, y + s * 0.7)],
                    paper, ink, 1.5, False))
        out.append(("line", [(x - s * 0.4, y + s * 0.7),
                             (x + s * 0.4, y - s * 0.1)], ink, 1.5, False))
    elif kind == "camp":
        out.append(("polygon",
                    [(x - s * 0.9, y + s * 0.7), (x, y - s * 0.85),
                     (x + s * 0.9, y + s * 0.7)],
                    "", ink, 1.6, False))
        out.append(("line", [(x, y - s * 0.85), (x, y + s * 0.7)], ink, 1.2, False))
    elif kind == "bridge":
        out.append(("line", [(x - s, y + s * 0.4), (x - s * 0.4, y - s * 0.5),
                             (x + s * 0.4, y - s * 0.5), (x + s, y + s * 0.4)],
                    ink, 1.8, False))
    elif kind == "inn":
        ring(s * 0.85, paper, 1.5)
        out.append(("line", [(x - s * 0.4, y - s * 0.3),
                             (x - s * 0.4, y + s * 0.35)], ink, 1.3, False))
        out.append(("line", [(x + s * 0.4, y - s * 0.3),
                             (x + s * 0.4, y + s * 0.35)], ink, 1.3, False))
        out.append(("line", [(x - s * 0.4, y + s * 0.05),
                             (x + s * 0.4, y + s * 0.05)], ink, 1.3, False))
    elif kind == "battle":
        out.append(("line", [(x - s, y - s), (x + s, y + s)], ink, 2.0, False))
        out.append(("line", [(x + s, y - s), (x - s, y + s)], ink, 2.0, False))
        out.append(("line", [(x - s * 1.15, y - s * 0.55),
                             (x - s * 0.55, y - s * 1.15)], ink, 1.5, False))
    elif kind == "treasure":
        out.append(("line", [(x - s * 1.1, y - s * 1.1),
                             (x + s * 1.1, y + s * 1.1)], accent, 3.0, False))
        out.append(("line", [(x + s * 1.1, y - s * 1.1),
                             (x - s * 1.1, y + s * 1.1)], accent, 3.0, False))
    elif kind == "danger":
        out.append(("polygon",
                    [(x, y - s * 1.1), (x + s, y + s * 0.75), (x - s, y + s * 0.75)],
                    paper, ink, 1.7, False))
        out.append(("line", [(x, y - s * 0.45), (x, y + s * 0.2)], ink, 1.8, False))
        out.append(("ellipse", x - s * 0.13, y + s * 0.38,
                    x + s * 0.13, y + s * 0.58, ink, ink, 1.0))
    elif kind == "portal":
        ring(s * 1.1, "", 1.8)
        ring(s * 0.62, "", 1.3)
        ring(s * 0.2, ink, 1.0)
    else:  # landmark
        out.append(("polygon",
                    [(x, y - s), (x + s * 0.32, y - s * 0.32),
                     (x + s, y), (x + s * 0.32, y + s * 0.32),
                     (x, y + s), (x - s * 0.32, y + s * 0.32),
                     (x - s, y), (x - s * 0.32, y - s * 0.32)],
                    paper, ink, 1.4, False))
    return out


def pin_label_anchor(pin: Pin, side: str = "") -> Tuple[float, float, str]:
    """Where a pin's label sits, and its text anchor."""
    s = max(3.0, float(pin.size))
    gap = s * 2.0
    side = (side or pin.label_side or "e").lower()
    if side == "w":
        return (pin.x - gap, pin.y, "e")
    if side == "n":
        return (pin.x, pin.y - gap, "s")
    if side == "s":
        return (pin.x, pin.y + gap, "n")
    if side == "ne":
        return (pin.x + gap * 0.72, pin.y - gap * 0.72, "w")
    if side == "nw":
        return (pin.x - gap * 0.72, pin.y - gap * 0.72, "e")
    if side == "se":
        return (pin.x + gap * 0.72, pin.y + gap * 0.72, "w")
    if side == "sw":
        return (pin.x - gap * 0.72, pin.y + gap * 0.72, "e")
    return (pin.x + gap, pin.y, "w")


# --------------------------------------------------------------------------
# Automatic label placement
#
# A generated map puts fifty place names on the page and they collide. Real
# cartographers solve this by moving a label around its point and, failing
# that, leaving it off. That is exactly what this does.
# --------------------------------------------------------------------------

Box = Tuple[float, float, float, float]

# Which labels win a fight. A capital's name matters more than a village's.
LABEL_PRIORITY = {
    "capital": 0, "city": 1, "port": 2, "castle": 3, "temple": 4,
    "town": 5, "dungeon": 6, "ruin": 7, "portal": 8, "landmark": 9,
    "mine": 10, "battle": 11, "treasure": 12, "bridge": 13, "cave": 14,
    "tower": 15, "inn": 16, "camp": 17, "danger": 18, "village": 19,
}

# Order in which alternative positions are tried around a point.
LABEL_SIDES = ("e", "w", "n", "s", "ne", "nw", "se", "sw")


def text_box(x: float, y: float, text: str, size: float, anchor: str,
             tracking: float = 0.0) -> Box:
    """
    Estimate a text bounding box without a font.

    0.52em average advance width is a good approximation for a serif face and
    keeps this dependency-free, which matters because the same estimate has to
    hold for the Tk canvas, Pillow and SVG.
    """
    if not text:
        return (x, y, x, y)
    width = len(text) * size * 0.52 + tracking * max(0, len(text) - 1)
    height = size * 1.15
    if anchor == "w":
        x0 = x
    elif anchor == "e":
        x0 = x - width
    else:
        x0 = x - width / 2.0
    if anchor == "n":
        y0 = y
    elif anchor == "s":
        y0 = y - height
    else:
        y0 = y - height / 2.0
    return (x0, y0, x0 + width, y0 + height)


def boxes_overlap(a: Box, b: Box, pad: float = 1.5) -> bool:
    return not (a[2] + pad < b[0] or b[2] + pad < a[0]
                or a[3] + pad < b[1] or b[3] + pad < a[1])


# --------------------------------------------------------------------------
# Scale
#
# One place decides how long the scale bar is and what it stands for. The
# renderer draws it, the label placer keeps names off it, and mapstory reads
# every distance from it - so a distance NovelForge reports is always the
# distance the picture shows. They used to disagree: the bar was drawn at 16%
# of the map's width (capped at 230 px) while mapstory assumed 20%, so every
# distance came out at 64-72% of what the caption said.
# --------------------------------------------------------------------------

_MILE = 1.0
_FOOT = _MILE / 5280.0

#: (singular, plural, other spellings, miles in one). Lower case; matching is
#: case-blind. Abbreviations stay as the writer typed them (`10 ft`, `3 AU`);
#: spelled-out words agree with the number (`1 mile`, `5 miles`).
UNIT_TABLE: Tuple[Tuple[str, str, Tuple[str, ...], float], ...] = (
    ("foot", "feet", ("ft", "foots"), _FOOT),
    ("yard", "yards", ("yd", "yds"), 3 * _FOOT),
    ("pace", "paces", (), 2.5 * _FOOT),
    ("metre", "metres", ("meter", "meters", "m"), 0.000621371),
    ("kilometre", "kilometres", ("kilometer", "kilometers", "km", "kms"),
     0.621371),
    ("mile", "miles", ("mi",), _MILE),
    ("league", "leagues", (), 3.0),
    ("day", "days", (), 20.0),                       # a day's march
    ("march", "marches", (), 20.0),
    ("parsec", "parsecs", ("pc",), 1.9174e13),
    ("astronomical unit", "astronomical units", ("au",), 92955807.0),
    ("light-year", "light-years",
     ("light year", "light years", "lightyear", "lightyears", "ly"), 5.8786e12),
)

#: Units too big for "about N miles" to help, or for a horse-speed check.
ASTRONOMICAL_UNITS = frozenset(
    {"parsec", "parsecs", "pc", "astronomical unit", "astronomical units", "au",
     "light-year", "light-years", "light year", "light years", "lightyear",
     "lightyears", "ly"})

_UNIT_BY_WORD: Dict[str, Tuple[str, str, Tuple[str, ...], float]] = {}
for _row in UNIT_TABLE:
    for _word in (_row[0], _row[1]) + _row[2]:
        _UNIT_BY_WORD[_word] = _row


def unit_info(word: str) -> Optional[Tuple[str, str, Tuple[str, ...], float]]:
    """The row of `UNIT_TABLE` a caption's unit word means, or None."""
    return _UNIT_BY_WORD.get(re.sub(r"\s+", " ", (word or "").strip().lower()))


_CAPTION_RE = re.compile(
    r"(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)"      # 1,500 or 12.5
    r"\s*([A-Za-z][A-Za-z\-]*)?(?:\s+([A-Za-z][A-Za-z\-]*))?")


def parse_scale_caption(text: str) -> Optional[Tuple[float, str]]:
    """
    (amount, unit word) from a caption such as "100 leagues" or "3 light years".

    The unit is "units" when the caption has a number but no word. None when it
    has no usable number at all ("a long walk"), which means the scale is
    unknown - never a guess.
    """
    match = _CAPTION_RE.search(text or "")
    if not match:
        return None
    amount = float(match.group(1).replace(",", ""))
    if amount <= 0:
        return None
    first, second = match.group(2), match.group(3)
    if first and second and unit_info(f"{first} {second}"):
        return amount, f"{first} {second}".lower()
    return amount, (first or "units").lower()


def nice_number(value: float, mode: str = "floor") -> float:
    """
    A round number: 1, 2 or 5 times a power of ten.

    "floor" is the largest one not above `value`, "ceil" the smallest not
    below it, and "round" the nearest on a logarithmic scale. A scale bar is
    always one of these - nobody wants a bar labelled 137 miles.
    """
    if not value or value <= 0 or math.isinf(value) or math.isnan(value):
        return 0.0
    exponent = math.floor(math.log10(value))
    base = value / 10.0 ** exponent
    if base >= 10.0:                       # log10 is not exact at the edges
        exponent, base = exponent + 1, base / 10.0
    elif base < 1.0:
        exponent, base = exponent - 1, base * 10.0
    steps = (1.0, 2.0, 5.0, 10.0)
    if mode == "ceil":
        pick = next((s for s in steps if s >= base * (1 - 1e-9)), 10.0)
    elif mode == "round":
        pick = min(steps, key=lambda s: abs(math.log(base / s)))
    else:
        pick = max((s for s in steps if s <= base * (1 + 1e-9)), default=1.0)
    return float(f"{pick:g}e{exponent}")


def default_bar_px(width: float) -> float:
    """
    How long the scale bar is drawn on a map that has not been given its own
    length: 16% of the width, but never more than 230 px. Every map made before
    `scale_px` existed was drawn this way, so it stays true to them.
    """
    return min(float(width) * 0.16, 230.0)


@dataclass(frozen=True)
class ScaleBar:
    """The scale bar as drawn: where, how long, and what its length means."""

    px: float = 0.0            # drawn length, in map pixels
    value: float = 0.0         # what that length stands for, in `unit`
    unit: str = ""             # the word after the number, lower case
    label: str = ""            # the caption printed above the bar
    x: float = 0.0             # left end
    y: float = 0.0             # top edge
    thickness: float = 0.0     # how tall the bar is
    area: Box = (0.0, 0.0, 0.0, 0.0)   # what it keeps clear of names

    @property
    def known(self) -> bool:
        """True when the caption holds a number, so distances can be worked out."""
        return self.value > 0 and self.px > 0

    @property
    def units_per_pixel(self) -> float:
        return self.value / self.px if self.known else 0.0


def scale_bar(gm: Any) -> ScaleBar:
    """
    The one description of a map's scale bar.

    The bar is `gm.scale_px` long when the map has set that (`apply_scale`
    does), otherwise `default_bar_px(width)`; the caption says what that length
    is worth. A map with no caption has no bar (an empty label). Anything that
    draws the bar, keeps names off it or turns pixels into miles goes through
    here, so they cannot disagree. Works on anything with the map's attributes.
    """
    width = float(getattr(gm, "width", 1600) or 1600)
    height = float(getattr(gm, "height", 1100) or 1100)
    caption = str(getattr(gm, "scale_text", "") or "").strip()
    custom = float(getattr(gm, "scale_px", 0.0) or 0.0)
    px = custom if custom > 0 else default_bar_px(width)
    px = max(4.0, min(px, width * 0.9))
    parsed = parse_scale_caption(caption)
    x, y = width * 0.045, height - height * 0.052
    thickness = max(7.0, height * 0.011)
    area = (x - 6, y - height * 0.03, x + px + 6, y + height * 0.02)
    if caption:
        words = text_box(x + px / 2.0, y - thickness * 1.5, caption,
                         max(9, int(height * 0.016)), "center")
        area = (min(area[0], words[0]), min(area[1], words[1]),
                max(area[2], words[2]), max(area[3], words[3]))
    return ScaleBar(px=px, value=parsed[0] if parsed else 0.0,
                    unit=parsed[1] if parsed else "", label=caption,
                    x=x, y=y, thickness=thickness, area=area)


def _plain_number(value: float) -> str:
    if value >= 1000 and abs(value - round(value)) < 1e-9:
        return f"{int(round(value)):,}"
    return f"{value:.6f}".rstrip("0").rstrip(".")


def scale_caption(value: float, unit: str) -> str:
    """"10 ft", "1 mile", "5 miles": a number and a unit that agree."""
    typed = (unit or "").strip()
    row = unit_info(typed)
    word = typed
    if row is not None and typed.lower() in (row[0], row[1]):
        word = row[0] if abs(value - 1.0) < 1e-9 else row[1]
    return f"{_plain_number(value)} {word}".strip()


def apply_scale(gm: Any, unit: str, per_cell: float,
                cell_px: Optional[float] = None) -> ScaleBar:
    """
    Give a map a scale: `per_cell` `unit`s to one grid square.

    Sets the caption, the bar's drawn length and the grid square together, so
    the caption, the bar and the grid cannot disagree: with `cell_px` the grid
    square is that many map pixels (rounded to a whole pixel), otherwise the
    map's `grid_size` is kept. The bar is the largest round number (1, 2 or 5
    times a power of ten) of `unit` that fits the space a bar normally gets, so
    "5 ft per square" on a 60 px grid gives a bar reading "10 ft", 120 px long.
    Returns the resulting `ScaleBar`.
    """
    per_cell = float(per_cell)
    if per_cell <= 0 or math.isnan(per_cell) or math.isinf(per_cell):
        raise ValueError("per_cell must be a positive number")
    if cell_px is not None:
        gm.grid_size = max(4, int(round(float(cell_px))))
    cell = max(1.0, float(gm.grid_size))
    units_per_px = per_cell / cell
    room = default_bar_px(float(gm.width))
    value = nice_number(room * units_per_px, "floor")
    gm.scale_text = scale_caption(value, unit)
    gm.scale_px = value / units_per_px
    return scale_bar(gm)


def cell_units(gm: Any) -> float:
    """How many of the scale's units one grid square is worth (0 if unknown)."""
    bar = scale_bar(gm)
    return bar.units_per_pixel * float(getattr(gm, "grid_size", 0) or 0)


def _label_fingerprint(gm: GameMap, visible: set) -> tuple:
    """
    Everything `layout_pin_labels` actually reads, rounded to a tenth of a
    map unit. Panning and zooming touch none of this, so a redraw triggered
    by either can reuse the previous placement instead of repeating an
    O(pins x placed) collision search - the part of a redraw that gets slow
    on a busy generated world.
    """
    return (
        gm.title_on_map, gm.name, gm.compass, gm.width, gm.height,
        gm.scale_text, gm.scale_px, gm.hide_colliding_labels, frozenset(visible),
        tuple(
            (l.layer, round(l.x, 1), round(l.y, 1), l.text, l.size, l.tracking)
            for l in gm.labels
        ),
        tuple(
            (s.layer, s.label, round(s.centroid()[0], 1), round(s.centroid()[1], 1))
            for s in gm.shapes if s.label.strip()
        ),
        tuple(
            (p.id, p.layer, round(p.x, 1), round(p.y, 1), p.size, p.label,
             p.label_side, p.kind)
            for p in gm.pins
        ),
    )


def _cached_label_placement(gm: GameMap, visible: set
                            ) -> Dict[str, Tuple[str, bool]]:
    fingerprint = _label_fingerprint(gm, visible)
    cached = gm._label_cache
    if cached is not None and cached[0] == fingerprint:
        return cached[1]
    placement = layout_pin_labels(gm, visible)
    gm._label_cache = (fingerprint, placement)
    return placement


def layout_pin_labels(gm: GameMap, visible: Optional[set] = None
                      ) -> Dict[str, Tuple[str, bool]]:
    """
    Choose a side for every pin label and decide which ones must be dropped.

    Returns {pin_id: (side, show)}. Reserved first, in order: the map title,
    compass and scale bar; then free-standing text labels; then every pin's own
    glyph, so a name never sits on top of a marker.
    """
    if visible is None:
        visible = gm.visible_layers()
    placed: List[Box] = []

    # Furniture.
    if gm.title_on_map and gm.name.strip():
        size = max(18, int(gm.height * 0.038))
        placed.append(text_box(gm.width / 2.0, gm.height * 0.062,
                               gm.name.upper(), size, "center", size * 0.14))
    if gm.compass:
        r = min(gm.width, gm.height) * 0.055
        cx, cy = gm.width - r * 2.1, gm.height - r * 2.1
        placed.append((cx - r * 1.5, cy - r * 1.8, cx + r * 1.5, cy + r * 1.5))
    bar = scale_bar(gm)
    if bar.label:
        placed.append(bar.area)

    # Free-standing labels (region and ocean names) always win.
    for label in gm.labels:
        if label.layer not in visible or not label.text.strip():
            continue
        placed.append(text_box(label.x, label.y, label.text, label.size,
                               "center", label.tracking))

    # Named areas drawn from a shape's label.
    for shape in gm.shapes:
        if shape.layer in visible and shape.label.strip():
            cx, cy = shape.centroid()
            placed.append(text_box(cx, cy, shape.label,
                                   max(11, int(gm.height * 0.019)),
                                   "center", 2.0))

    pins = [p for p in gm.pins if p.layer in visible]
    # Reserve every glyph before placing any text.
    for pin in pins:
        s = max(3.0, float(pin.size)) * 1.7
        placed.append((pin.x - s, pin.y - s, pin.x + s, pin.y + s))

    out: Dict[str, Tuple[str, bool]] = {}
    ordered = sorted(
        pins,
        key=lambda p: (LABEL_PRIORITY.get(p.kind, 50), -float(p.size),
                       p.label.lower()),
    )
    size = max(9, int(gm.height * 0.0155))
    for pin in ordered:
        if not pin.label.strip():
            out[pin.id] = (pin.label_side or "e", False)
            continue
        # Try the author's chosen side first, then the rest.
        preferred = (pin.label_side or "e").lower()
        candidates = [preferred] + [s for s in LABEL_SIDES if s != preferred]
        chosen: Optional[str] = None
        for side in candidates:
            lx, ly, anchor = pin_label_anchor(pin, side)
            box = text_box(lx, ly, pin.label, size, anchor)
            # Stay on the page.
            if box[0] < 4 or box[1] < 4 or box[2] > gm.width - 4 \
                    or box[3] > gm.height - 4:
                continue
            if any(boxes_overlap(box, other) for other in placed):
                continue
            chosen = side
            placed.append(box)
            break
        if chosen is None:
            # Nowhere to put it. Hiding one village name is far better than a
            # page of names printed on top of each other.
            out[pin.id] = (preferred, not gm.hide_colliding_labels)
            if not gm.hide_colliding_labels:
                lx, ly, anchor = pin_label_anchor(pin, preferred)
                placed.append(text_box(lx, ly, pin.label, size, anchor))
        else:
            out[pin.id] = (chosen, True)
    return out


# -- terrain decoration ----------------------------------------------------


def _luma(colour: str) -> float:
    r, g, b = _rgb(colour)
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255.0


def _relief_tones(land: str, ink: str) -> Tuple[str, str, str]:
    """(lit, shaded, snow) for terrain stamps, whatever the palette. Light comes from the left."""
    if _luma(land) > 0.5:
        return (_mix(land, "#ffffff", 0.30), _mix(land, ink, 0.34),
                _mix(land, "#ffffff", 0.85))
    return (_mix(land, "#ffffff", 0.16), _mix(land, "#000000", 0.38),
            _mix(land, "#ffffff", 0.55))


def _peak(x: float, y: float, size: float, ink: str, lit: str, shade: str,
          snow: str) -> List[tuple]:
    """One mountain: a lit left face, a shaded right face, an outline, snow on the big ones."""
    half = size * 0.8
    apex = (x + size * 0.05, y - size)
    left = (x - half, y + size * 0.38)
    right = (x + half, y + size * 0.42)
    foot = (x + size * 0.12, y + size * 0.4)
    out: List[tuple] = [
        ("polygon", [left, apex, foot], lit, "", 0.0, False),
        ("polygon", [apex, right, foot], shade, "", 0.0, False),
    ]
    if size >= 15.0:
        out.append(("polygon",
                    [apex, (apex[0] - half * 0.34, apex[1] + size * 0.30),
                     (apex[0] - half * 0.10, apex[1] + size * 0.22),
                     (apex[0] + half * 0.06, apex[1] + size * 0.34),
                     (apex[0] + half * 0.36, apex[1] + size * 0.29)],
                    snow, "", 0.0, False))
    out.append(("line", [apex, (x + size * 0.02, y + size * 0.1), foot],
                ink, 0.8, False))
    out.append(("line", [left, apex, right], ink, 1.25, False))
    return out


def _peaks(points: Sequence[Point], ink: str, land: str,
           spacing: float = 22.0, height: float = 17.0,
           seed: int = 7) -> List[tuple]:
    """
    A mountain range along a ridge line.

    Each spot on the ridge gets one to three peaks: a main one and smaller
    companions to either side, so the range has depth instead of being a row of
    identical tents. Peaks grow toward the middle of the range and taper at its
    ends, and are drawn back to front so nearer ones overlap farther ones.
    """
    rng = random.Random(seed)
    samples = _walk_path(points, spacing)
    lit, shade, snow = _relief_tones(land, ink)
    count = len(samples)
    stamps: List[Tuple[float, float, float]] = []
    for i, ((px, py), angle) in enumerate(samples):
        t = i / (count - 1) if count > 1 else 0.5
        envelope = 0.62 + 0.62 * math.sin(math.pi * t)
        nx, ny = -math.sin(angle), math.cos(angle)
        companions = (rng.random() < 0.6) + (rng.random() < 0.3)
        for k in range(1 + companions):
            side = 0.0 if k == 0 else rng.choice((-1.0, 1.0)) * rng.uniform(0.8, 1.5) * height
            slide = rng.uniform(-0.45, 0.45) * spacing
            size = height * envelope * (rng.uniform(0.85, 1.35) if k == 0
                                        else rng.uniform(0.5, 0.9))
            stamps.append((px + nx * side + math.cos(angle) * slide,
                           py + ny * side + math.sin(angle) * slide, size))
    stamps.sort(key=lambda stamp: stamp[1])
    out: List[tuple] = []
    for x, y, size in stamps:
        out.extend(_peak(x, y, size, ink, lit, shade, snow))
    return out


def _hills(points: Sequence[Point], ink: str, land: str,
           spacing: float = 19.0, seed: int = 5) -> List[tuple]:
    """Rolling hills: low shaded domes, in loose clusters along a line."""
    rng = random.Random(seed)
    lit, shade, _snow = _relief_tones(land, ink)
    stamps: List[Tuple[float, float, float]] = []
    for (px, py), angle in _walk_path(points, spacing):
        nx, ny = -math.sin(angle), math.cos(angle)
        for k in range(1 + (rng.random() < 0.35)):
            side = 0.0 if k == 0 else rng.choice((-1.0, 1.0)) * rng.uniform(6.0, 11.0)
            stamps.append((px + nx * side + rng.uniform(-4.0, 4.0),
                           py + ny * side + rng.uniform(-3.0, 3.0),
                           rng.uniform(6.5, 10.5)))
    stamps.sort(key=lambda stamp: stamp[1])
    out: List[tuple] = []
    for x, y, r in stamps:
        dome = [(x + r * math.cos(math.pi * k / 8.0),
                 y - r * 0.72 * math.sin(math.pi * k / 8.0)) for k in range(9)]
        out.append(("polygon", dome, lit, "", 0.0, False))
        out.append(("polygon", [dome[0], dome[1], dome[2], dome[3], (x + r * 0.05, y)],
                    shade, "", 0.0, False))
        out.append(("line", dome, ink, 1.1, False))
    return out


def _trees(shape: Shape, ink: str, canopy: str, seed: int,
           spacing: float = 21.0, cap: int = 650) -> List[tuple]:
    """
    Trees scattered inside a forest.

    Clumped rather than evenly spread: a smooth pseudo-random field gates whole
    neighbourhoods on or off, so the wood has clearings and dense stands instead
    of looking like tiled wallpaper. Sizes vary, roughly a third are conifers,
    and they are drawn back to front. The count is capped so a continent-sized
    forest does not turn into ten thousand canvas items.
    """
    rng = random.Random(seed ^ _stable_seed(shape))
    spots = _scatter_in_bounds(shape, spacing, seed)
    if len(spots) > cap:
        stride = len(spots) / float(cap)
        spots = [spots[int(i * stride)] for i in range(cap)]
    trees: List[Tuple[float, float, float, bool]] = []
    for px, py in spots:
        clump = (math.sin(px * 0.021 + seed * 0.7) +
                 math.cos(py * 0.019 - seed * 0.4))
        if clump < -0.55 or rng.random() < 0.15:
            continue
        r = 3.6 * rng.uniform(0.8, 1.45)
        trees.append((px + rng.uniform(-spacing * 0.2, spacing * 0.2),
                      py + rng.uniform(-spacing * 0.2, spacing * 0.2),
                      r, rng.random() < 0.34))
    trees.sort(key=lambda tree: tree[1])
    out: List[tuple] = []
    for jx, jy, r, conifer in trees:
        out.append(("line", [(jx, jy + r * 0.45), (jx, jy + r * 1.25)],
                    ink, 1.0, False))
        if conifer:
            out.append(("polygon",
                        [(jx - r * 0.85, jy + r * 0.75), (jx, jy - r * 1.6),
                         (jx + r * 0.85, jy + r * 0.75)], canopy, ink, 0.9, False))
        else:
            out.append(("ellipse", jx - r, jy - r * 1.15, jx + r, jy + r * 0.6,
                        canopy, ink, 0.9))
    return out


def _dots(shape: Shape, ink: str, seed: int,
          spacing: float = 20.0) -> List[tuple]:
    out: List[tuple] = []
    for px, py in _scatter_in_bounds(shape, spacing, seed + 31):
        out.append(("ellipse", px - 1.0, py - 1.0, px + 1.0, py + 1.0,
                    ink, ink, 1.0))
    return out


def _squiggles(shape: Shape, ink: str, seed: int,
               spacing: float = 26.0) -> List[tuple]:
    out: List[tuple] = []
    for px, py in _scatter_in_bounds(shape, spacing, seed + 71):
        out.append(("line", [(px - 6, py), (px - 2, py - 2.5), (px + 2, py),
                             (px + 6, py - 2.5)], ink, 1.2, False))
    return out


def _tapered_river(points: Sequence[Point], colour: str,
                   width: float, bank: str = "") -> List[tuple]:
    """A river drawn in three passes so it thickens toward its mouth, with banks."""
    pts = list(points)
    if len(pts) < 2:
        return []
    third = max(2, len(pts) // 3)
    passes = ((pts, max(1.0, width * 0.6)),
              (pts[third:], max(1.2, width * 0.85)),
              (pts[third * 2:], max(1.5, width * 1.15)))
    out: List[tuple] = []
    if bank:
        for run, w in passes:
            out.append(("line", run, bank, w + 1.2, False))
    for run, w in passes:
        out.append(("line", run, colour, w, False))
    return out


def _offset_ring(points: Sequence[Point], distance: float) -> List[Point]:
    """
    A closed ring pushed outward by `distance`.

    Each vertex moves along the bisector of its two edges. Where that would
    fold the outline back on itself (a bay narrower than the distance) the
    vertex is skipped, so the result is always a clean, if simplified, ring -
    which is all a decorative ripple needs. Returns [] for a degenerate ring.
    """
    ring = list(points)
    if len(ring) > 1 and ring[0] == ring[-1]:
        ring.pop()
    n = len(ring)
    if n < 4:
        return []
    area = sum(ring[i][0] * ring[(i + 1) % n][1] - ring[(i + 1) % n][0] * ring[i][1]
               for i in range(n))
    turn = 1.0 if area > 0 else -1.0          # a clockwise ring's outward normal is (dy, -dx)

    def unit_normal(a: Point, b: Point) -> Tuple[float, float]:
        dx, dy = b[0] - a[0], b[1] - a[1]
        length = math.hypot(dx, dy) or 1.0
        return (dy / length * turn, -dx / length * turn)

    moved: List[Tuple[Point, Point]] = []
    for i in range(n):
        prev, here, nxt = ring[i - 1], ring[i], ring[(i + 1) % n]
        n1, n2 = unit_normal(prev, here), unit_normal(here, nxt)
        bx, by = n1[0] + n2[0], n1[1] + n2[1]
        scale = distance / max(0.35, 1.0 + n1[0] * n2[0] + n1[1] * n2[1])
        moved.append(((here[0] + bx * scale, here[1] + by * scale), here))
    kept = [moved[0]]
    for point, original in moved[1:]:
        last, last_original = kept[-1]
        forward = ((point[0] - last[0]) * (original[0] - last_original[0])
                   + (point[1] - last[1]) * (original[1] - last_original[1]))
        if forward > 0.0:
            kept.append((point, original))
    out = [pair[0] for pair in kept]
    return _smooth(out + [out[0]], 2) if len(out) >= 4 else []


def _shore_primitives(pts: Sequence[Point], width: float, palette: Dict[str, Any],
                      reach: float) -> List[tuple]:
    """
    The sea around a coast: pale shallows that fade out, then two ripple lines.

    The shallows are stroked copies of the coastline, widest first, each a
    little narrower and a little paler than the last. The ripples are thin
    rings pushed out from the coast. `build_primitives` paints every coast's
    water before any land, so a ripple that strays across a neighbouring
    island is simply covered by it.
    """
    sea = palette["paper"]
    shallow = (palette.get("halo") or [_mix(sea, "#ffffff", 0.4)])[-1]
    ink = palette["ink_light"]
    out: List[tuple] = []
    for distance, tone, dashed in ((reach * 2.35, 0.74, True),
                                   (reach * 1.7, 0.6, False)):
        ring = _offset_ring(pts, distance)
        if ring:
            out.append(("line", ring, _mix(ink, sea, tone), 1.0, dashed))
    steps = 5
    for i in range(steps, 0, -1):
        out.append(("line", pts, _mix(shallow, sea, (i - 1) / float(steps)),
                    width + 2 * reach * i / steps, False))
    return out


# -- background and furniture ----------------------------------------------


def _grid_primitives(gm: GameMap) -> List[tuple]:
    palette = gm.palette()
    # Blended toward the paper: a grid is a reading aid, not a feature, and at
    # full strength it competes with the terrain for attention.
    colour = _mix(palette["grid"], palette["paper"], 0.55)
    step = max(20, int(gm.grid_size))
    out: List[tuple] = []
    if gm.grid == "square":
        x = 0
        while x <= gm.width:
            out.append(("line", [(x, 0), (x, gm.height)], colour, 0.7, False))
            x += step
        y = 0
        while y <= gm.height:
            out.append(("line", [(0, y), (gm.width, y)], colour, 0.7, False))
            y += step
    elif gm.grid == "hex":
        r = step / 2.0
        h = r * math.sqrt(3)
        row = 0
        y = 0.0
        while y <= gm.height + h:
            offset = 0.0 if row % 2 == 0 else r * 1.5
            x = offset
            while x <= gm.width + r * 2:
                pts = [
                    (x + r * math.cos(math.pi / 3 * i),
                     y + r * math.sin(math.pi / 3 * i))
                    for i in range(6)
                ]
                out.append(("polygon", pts, "", colour, 0.7, False))
                x += r * 3
            y += h / 2
            row += 1
            if row > 400:            # guard against a pathological grid_size
                break
    return out


def _compass_primitives(gm: GameMap) -> List[tuple]:
    palette = gm.palette()
    ink, paper = palette["ink"], palette["paper"]
    r = min(gm.width, gm.height) * 0.055
    cx = gm.width - r * 2.1
    cy = gm.height - r * 2.1
    out: List[tuple] = [
        ("ellipse", cx - r, cy - r, cx + r, cy + r, "", ink, 1.4),
        ("ellipse", cx - r * 0.75, cy - r * 0.75, cx + r * 0.75, cy + r * 0.75,
         "", palette["ink_light"], 0.9),
    ]
    for i in range(4):
        a = -math.pi / 2 + i * math.pi / 2
        tip = (cx + math.cos(a) * r * 0.95, cy + math.sin(a) * r * 0.95)
        left = (cx + math.cos(a + math.pi / 2) * r * 0.17,
                cy + math.sin(a + math.pi / 2) * r * 0.17)
        right = (cx + math.cos(a - math.pi / 2) * r * 0.17,
                 cy + math.sin(a - math.pi / 2) * r * 0.17)
        out.append(("polygon", [tip, left, (cx, cy), right],
                    ink if i == 0 else paper, ink, 1.1, False))
    for i in range(4):
        a = -math.pi / 4 + i * math.pi / 2
        out.append(("line",
                    [(cx + math.cos(a) * r * 0.2, cy + math.sin(a) * r * 0.2),
                     (cx + math.cos(a) * r * 0.62, cy + math.sin(a) * r * 0.62)],
                    palette["ink_light"], 0.9, False))
    out.append(("text", cx, cy - r * 1.42, "N", int(r * 0.62), ink,
                "center", False, True, 0.0))
    return out


def _scale_primitives(gm: GameMap) -> List[tuple]:
    scale = scale_bar(gm)
    if not scale.label:
        return []
    palette = gm.palette()
    ink, paper = palette["ink"], palette["paper"]
    bar, x0, y0, height = scale.px, scale.x, scale.y, scale.thickness
    out: List[tuple] = []
    segments = 4
    for i in range(segments):
        sx = x0 + bar / segments * i
        out.append(("polygon",
                    [(sx, y0), (sx + bar / segments, y0),
                     (sx + bar / segments, y0 + height), (sx, y0 + height)],
                    ink if i % 2 == 0 else paper, ink, 1.1, False))
    out.append(("text", x0 + bar / 2, y0 - height * 1.5, scale.label,
                max(9, int(gm.height * 0.016)), ink, "center", True, False, 0.0))
    return out


def _border_primitives(gm: GameMap) -> List[tuple]:
    palette = gm.palette()
    ink, light = palette["ink"], palette["ink_light"]
    inset = max(10.0, min(gm.width, gm.height) * 0.018)
    out: List[tuple] = [
        ("polygon", [(inset, inset), (gm.width - inset, inset),
                     (gm.width - inset, gm.height - inset),
                     (inset, gm.height - inset)], "", ink, 2.4, False),
    ]
    inner = inset * 1.75
    out.append(("polygon", [(inner, inner), (gm.width - inner, inner),
                            (gm.width - inner, gm.height - inner),
                            (inner, gm.height - inner)], "", light, 1.0, False))
    return out


def _title_primitives(gm: GameMap) -> List[tuple]:
    if not gm.title_on_map or not gm.name.strip():
        return []
    palette = gm.palette()
    size = max(18, int(gm.height * 0.038))
    return [
        ("text", gm.width / 2.0, gm.height * 0.062, gm.name.upper(),
         size, palette["ink"], "center", False, True, size * 0.14),
    ]


# -- the main builder ------------------------------------------------------


BIOMES = ("forest", "desert", "swamp", "ice")


def _points_key(points: Sequence[Point]) -> int:
    return hash(tuple(map(tuple, points)))


def _primitive_key(gm: GameMap, furniture: bool, visible: set) -> tuple:
    """
    Everything the display list depends on. Views (zoom, pan, selection) are
    not in it. The edition is not either, on purpose: it matters only through
    which layers it leaves visible, and that set is in here.
    """
    return (
        gm.style, gm.width, gm.height, gm.grid, gm.grid_size, gm.scale_text,
        gm.scale_px, gm.compass, gm.border, gm.title_on_map, gm.name, gm.seed,
        gm.hide_colliding_labels, gm.auto_place_labels, furniture,
        frozenset(visible),
        tuple((s.id, s.kind, s.layer, s.fill, s.outline, s.width, s.label,
               s.closed, _points_key(s.points)) for s in gm.shapes),
        tuple((p.id, p.layer, p.x, p.y, p.kind, p.label, p.label_side, p.size)
              for p in gm.pins),
        tuple((l.id, l.layer, l.x, l.y, l.text, l.size, l.color, l.italic,
               l.bold, l.tracking) for l in gm.labels),
    )


def build_primitives(gm: GameMap, include_furniture: bool = True,
                     edition: str = "author") -> List[tuple]:
    """
    Turn a map into an ordered display list. This is the single source of truth.

    The result is remembered until the map itself changes, and each shape's
    share of it is remembered until that shape does: zooming, panning and
    selecting rebuild nothing, and dragging one vertex re-scatters the trees of
    one forest, not all of them.

    `edition="reader"` leaves out every author-only layer (see `EDITIONS`).
    """
    key = _primitive_key(gm, include_furniture, gm.visible_layers(edition))
    cached = gm._prim_cache
    if cached is not None and cached[0] == key:
        return list(cached[1])
    prims = _build_primitives(gm, include_furniture, edition)
    gm._prim_cache = (key, prims)
    return list(prims)


def _shape_shore(gm: GameMap, shape: Shape, palette: Dict[str, Any]) -> List[tuple]:
    pts = _smooth(shape.points, 2 if len(shape.points) < 240 else 1)
    if pts and pts[0] != pts[-1]:
        pts = pts + [pts[0]]
    reach = max(10.0, min(26.0, gm.height * 0.018))
    return _shore_primitives(pts, float(shape.width or 2.0), palette, reach)


def _shape_body(gm: GameMap, shape: Shape, palette: Dict[str, Any]) -> List[tuple]:
    """Everything one shape draws: fill, outline, decoration and its own label."""
    ink, light, paper = palette["ink"], palette["ink_light"], palette["paper"]
    terrain: Dict[str, str] = palette.get("terrain", {}) or {}
    land_tone = terrain.get("land") or paper
    spec = TERRAIN.get(shape.kind, TERRAIN["land"])
    closed = shape.closed and spec["closed"] and len(shape.points) >= 3
    pts = _smooth(shape.points, 2 if len(shape.points) < 240 else 1)
    if closed and pts[0] != pts[-1]:
        pts = pts + [pts[0]]

    fill = shape.fill or terrain.get(shape.kind) or (spec["fill"] or "")
    outline = shape.outline or ink
    if shape.kind == "water" and not shape.outline:
        outline = _mix(ink, fill or paper, 0.5)         # a lake's edge is quieter than a coast
    width = float(shape.width or 2.0)
    decor = spec.get("decor")
    out: List[tuple] = []

    if closed:
        if shape.kind in BIOMES and fill and not shape.outline:
            # A forest or a desert has no inked edge on a real map: it fades.
            out.append(("line", pts, _mix(fill, land_tone, 0.55), 7.0, False))
            out.append(("polygon", pts, fill, "", 0.0, False))
            out.append(("line", pts, _mix(fill, ink, 0.22), 0.9, False))
        else:
            out.append(("polygon", pts, fill, outline if fill else "",
                        width if fill else 0.0, False))
            if fill and shape.kind != "region":
                out.append(("line", pts, outline, width, False))
            elif shape.kind == "region":
                out.append(("line", pts, shape.outline or light,
                            max(1.2, width), True))
    else:
        if decor == "river":
            # Rivers need to read at a glance against a busy land fill, so
            # they get the deeper water tone, a bank and a little more weight.
            core = shape.fill or palette["water_deep"]
            out.extend(_tapered_river(pts, core, width + 1.0,
                                      bank=_mix(core, ink, 0.30)))
        elif decor == "road":
            out.append(("line", pts, shape.outline or light,
                        max(1.4, width), True))
        elif decor == "route":
            out.append(("line", pts, shape.outline or ink,
                        max(1.6, width), True))
            if len(pts) >= 2:
                out.extend(_arrow_head(pts, shape.outline or ink,
                                       max(1.6, width)))
        elif decor == "wall":
            out.append(("line", pts, outline, max(2.4, width + 1.2), False))
            for (px, py), angle in _walk_path(pts, 14.0):
                nx = math.cos(angle + math.pi / 2) * 3.4
                ny = math.sin(angle + math.pi / 2) * 3.4
                out.append(("line", [(px - nx, py - ny), (px + nx, py + ny)],
                            outline, 1.2, False))
        elif decor not in ("peaks", "hills"):
            out.append(("line", pts, outline, width, False))

    seed = gm.seed ^ _stable_seed(shape)
    if decor == "peaks":
        out.extend(_peaks(pts, ink, land_tone, seed=seed))
    elif decor == "hills":
        out.extend(_hills(pts, ink, land_tone, seed=seed))
    elif decor == "trees" and closed:
        out.extend(_trees(shape, ink, _mix(fill or land_tone, ink, 0.34), gm.seed))
    elif decor == "dots" and closed:
        out.extend(_dots(shape, light, gm.seed))
    elif decor == "squiggles" and closed:
        out.extend(_squiggles(shape, light, gm.seed))

    if shape.label.strip():
        cx, cy = shape.centroid()
        out.append(("text", cx, cy, shape.label,
                    max(11, int(gm.height * 0.019)), light, "center",
                    True, False, 2.0, _mix(land_tone, paper, 0.25)))
    return out


def _build_primitives(gm: GameMap, include_furniture: bool,
                      edition: str = "author") -> List[tuple]:
    palette = gm.palette()
    ink = palette["ink"]
    paper = palette["paper"]
    terrain: Dict[str, str] = palette.get("terrain", {}) or {}
    land_tone = terrain.get("land") or paper
    halo = _mix(land_tone, paper, 0.25)
    visible = gm.visible_layers(edition)
    out: List[tuple] = []

    if gm.grid != "none":
        out.extend(_grid_primitives(gm))

    ordered = sorted(
        (s for s in gm.shapes if s.layer in visible and len(s.points) >= 2),
        key=lambda s: PAINT_ORDER.get(s.kind, 50),
    )

    live: set = set()
    cache = gm._shape_cache

    def remembered(shape: Shape, part: str, build) -> List[tuple]:
        key = (part, gm.style, gm.seed, gm.height, shape.kind, shape.fill,
               shape.outline, shape.width, shape.label, shape.closed,
               _points_key(shape.points))
        slot = (shape.id, part)
        live.add(slot)
        entry = cache.get(slot)
        if entry is not None and entry[0] == key:
            return entry[1]
        value = build(shape, palette) if part == "shore" else build(gm, shape, palette)
        cache[slot] = (key, value)
        return value

    def has_coast(shape: Shape) -> bool:
        spec = TERRAIN.get(shape.kind, TERRAIN["land"])
        return bool(spec.get("halo") and shape.closed and spec["closed"]
                    and len(shape.points) >= 3)

    shore_done = False
    for shape in ordered:
        if not shore_done and has_coast(shape):
            # Every coast's water goes down before any land, so the rings round
            # one island never paint over another. Water shapes drawn earlier
            # (a hand-drawn sea) still lie beneath them.
            for other in ordered:
                if has_coast(other):
                    out.extend(remembered(
                        other, "shore",
                        lambda sh, pal: _shape_shore(gm, sh, pal)))
            shore_done = True
        out.extend(remembered(shape, "body", _shape_body))
    for slot in [k for k in cache if k not in live]:
        del cache[slot]

    accent = "#8a2f22" if gm.style in ("parchment", "treasure") else ink
    placement = _cached_label_placement(gm, visible) if gm.auto_place_labels else {}
    label_size = max(9, int(gm.height * 0.0155))
    for pin in gm.pins:
        if pin.layer not in visible:
            continue
        out.extend(pin_primitives(pin, ink, paper, accent))
        if not pin.label.strip():
            continue
        side, show = placement.get(pin.id, (pin.label_side or "e", True))
        if not show:
            continue
        lx, ly, anchor = pin_label_anchor(pin, side)
        out.append(("text", lx, ly, pin.label, label_size, ink, anchor,
                    False, pin.kind in ("capital", "city"), 0.0, halo))

    for label in gm.labels:
        if label.layer not in visible:
            continue
        if not label.text.strip():
            continue
        out.append(("text", label.x, label.y, label.text, int(label.size),
                    label.color or ink, "center", label.italic, label.bold,
                    float(label.tracking), _mix(land_tone, paper, 0.5)))

    if include_furniture:
        if gm.border:
            out.extend(_border_primitives(gm))
        if gm.compass:
            out.extend(_compass_primitives(gm))
        out.extend(_scale_primitives(gm))
        out.extend(_title_primitives(gm))

    return out


def text_parts(prim: tuple) -> tuple:
    """A text primitive as (x, y, text, size, colour, anchor, italic, bold, tracking, halo)."""
    halo = prim[10] if len(prim) > 10 else ""
    return (prim[1], prim[2], prim[3], prim[4], prim[5], prim[6], prim[7],
            prim[8], prim[9], halo)


def _arrow_head(points: Sequence[Point], colour: str,
                width: float) -> List[tuple]:
    if len(points) < 2:
        return []
    (x0, y0), (x1, y1) = points[-2], points[-1]
    angle = math.atan2(y1 - y0, x1 - x0)
    size = max(7.0, width * 4.0)
    return [
        ("line", [(x1 - math.cos(angle - 0.42) * size,
                   y1 - math.sin(angle - 0.42) * size), (x1, y1)],
         colour, width, False),
        ("line", [(x1 - math.cos(angle + 0.42) * size,
                   y1 - math.sin(angle + 0.42) * size), (x1, y1)],
         colour, width, False),
    ]


# ==========================================================================
# PNG export (Pillow)
# ==========================================================================


def _font_folders() -> List[Path]:
    """
    Where to look for a real font file.

    Hardcoding C:/Windows/Fonts breaks on any machine where Windows is not on
    C:, and on the per-user font folder that installing a font without admin
    rights writes to. This is a portable tool, so it asks the system.
    """
    folders: List[Path] = []
    windir = os.environ.get("WINDIR") or os.environ.get("SystemRoot")
    if windir:
        folders.append(Path(windir) / "Fonts")
    local = os.environ.get("LOCALAPPDATA")
    if local:
        folders.append(Path(local) / "Microsoft" / "Windows" / "Fonts")
    # Not Windows: the usual places, so exported maps still get real type.
    folders += [Path("/usr/share/fonts"), Path("/usr/local/share/fonts"),
                Path.home() / ".fonts", Path("/Library/Fonts"),
                Path("/System/Library/Fonts")]
    return [f for f in folders if f.is_dir()]


def _find_font(italic: bool = False, bold: bool = False):
    from PIL import ImageFont

    if bold and italic:
        names = ["georgiaz.ttf", "timesbi.ttf", "constanz.ttf", "arialbi.ttf"]
    elif bold:
        names = ["georgiab.ttf", "timesbd.ttf", "constanb.ttf", "arialbd.ttf"]
    elif italic:
        names = ["georgiai.ttf", "timesi.ttf", "constani.ttf", "ariali.ttf"]
    else:
        names = ["georgia.ttf", "times.ttf", "constan.ttf", "arial.ttf",
                 "DejaVuSerif.ttf", "LiberationSerif-Regular.ttf"]
    for folder in _font_folders():
        for name in names:
            candidate = folder / name
            if candidate.exists():
                return str(candidate)
    return None


_FONT_CACHE: Dict[tuple, Any] = {}


def _font(size: int, italic: bool, bold: bool):
    from PIL import ImageFont

    key = (max(6, int(size)), bool(italic), bool(bold))
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    path = _find_font(italic, bold)
    try:
        font = ImageFont.truetype(path, key[0]) if path else ImageFont.load_default()
    except (OSError, ValueError):
        font = ImageFont.load_default()
    _FONT_CACHE[key] = font
    return font


def _parchment_background(gm: GameMap, scale: float):
    """Aged-paper texture: mottling plus a vignette. Deterministic per seed."""
    from PIL import Image, ImageDraw, ImageFilter

    palette = gm.palette()
    width = max(1, int(gm.width * scale))
    height = max(1, int(gm.height * scale))
    base = Image.new("RGB", (width, height), palette["paper"])

    if not palette.get("texture"):
        return base

    rng = random.Random(gm.seed or 7)
    blot = Image.new("L", (max(1, width // 6), max(1, height // 6)), 0)
    blot_draw = ImageDraw.Draw(blot)
    for _ in range(260):
        bx = rng.randint(0, blot.width)
        by = rng.randint(0, blot.height)
        br = rng.randint(2, max(3, blot.width // 14))
        blot_draw.ellipse([bx - br, by - br, bx + br, by + br],
                          fill=rng.randint(18, 78))
    blot = blot.filter(ImageFilter.GaussianBlur(radius=blot.width / 42.0))
    blot = blot.resize((width, height), Image.BILINEAR)
    dark = Image.new("RGB", (width, height), palette["paper_dark"])
    base = Image.composite(dark, base, blot.point(lambda v: min(255, v * 2)))

    if palette.get("vignette"):
        mask = Image.new("L", (width, height), 0)
        mask_draw = ImageDraw.Draw(mask)
        inset = int(min(width, height) * 0.06)
        mask_draw.rectangle([inset, inset, width - inset, height - inset],
                           fill=255)
        mask = mask.filter(
            ImageFilter.GaussianBlur(radius=max(6, min(width, height) * 0.05))
        )
        edge = Image.new("RGB", (width, height), palette["paper_dark"])
        base = Image.composite(base, edge, mask)

    return base


def _paint(draw, prims: Sequence[tuple], effective: float,
           min_stroke: int = 1) -> None:
    """
    Draw a display list with Pillow.

    `effective` is pixels per map unit. `min_stroke` is the thinnest line, in
    drawn pixels, that any outline may be - print export raises it so no line is
    finer than a printer can hold.
    """

    def sc(points: Iterable[Point]) -> List[Tuple[float, float]]:
        return [(p[0] * effective, p[1] * effective) for p in points]

    for prim in prims:
        head = prim[0]
        try:
            if head == "polygon":
                _kind, points, fill, outline, width, dash = prim
                pts = sc(points)
                if len(pts) < 3:
                    continue
                if fill:
                    draw.polygon(pts, fill=fill)
                if outline and width:
                    draw.line(pts + [pts[0]], fill=outline,
                              width=max(min_stroke, int(round(width * effective))),
                              joint="curve")
            elif head == "line":
                _kind, points, colour, width, dash = prim
                pts = sc(points)
                if len(pts) < 2 or not colour:
                    continue
                stroke = max(min_stroke, int(round(width * effective)))
                if dash:
                    for seg in _dash_segments(pts, 9.0 * effective,
                                              6.0 * effective):
                        if len(seg) >= 2:
                            draw.line(seg, fill=colour, width=stroke,
                                      joint="curve")
                else:
                    draw.line(pts, fill=colour, width=stroke, joint="curve")
            elif head == "ellipse":
                _kind, x0, y0, x1, y1, fill, outline, width = prim
                box = [x0 * effective, y0 * effective,
                       x1 * effective, y1 * effective]
                if box[2] < box[0]:
                    box[0], box[2] = box[2], box[0]
                if box[3] < box[1]:
                    box[1], box[3] = box[3], box[1]
                if box[2] - box[0] < 1 or box[3] - box[1] < 1:
                    continue
                draw.ellipse(box, fill=fill or None, outline=outline or None,
                             width=max(min_stroke, int(round(width * effective))))
            elif head == "text":
                (x, y, text, size, colour, anchor, italic, bold, tracking,
                 halo) = text_parts(prim)
                font = _font(int(size * effective), italic, bold)
                anchor_map = {"center": "mm", "w": "lm", "e": "rm",
                              "n": "ma", "s": "md"}
                stroke = max(1, int(round(size * 0.13 * effective))) if halo else 0
                if tracking and abs(tracking) > 0.4:
                    _draw_tracked(draw, x * effective, y * effective, text,
                                  font, colour, tracking * effective,
                                  anchor_map.get(anchor, "mm"), halo, stroke)
                else:
                    draw.text((x * effective, y * effective), text, font=font,
                              fill=colour, anchor=anchor_map.get(anchor, "mm"),
                              stroke_width=stroke,
                              stroke_fill=halo or None)
        except (ValueError, TypeError, OSError):
            # One malformed primitive must not abandon the whole export.
            continue


def _render_image(gm: GameMap, scale: float, supersample: int = 2,
                  edition: str = "author", min_stroke: int = 1,
                  size: Optional[Tuple[int, int]] = None):
    """
    The map as a Pillow RGB image, `scale` pixels per map unit.

    Drawn at `supersample` times the size and shrunk, which is how the lines
    come out smooth. `size` fixes the final pixel size (otherwise the map's size
    times `scale`, truncated); `min_stroke` is in final pixels.
    """
    from PIL import Image, ImageDraw

    ss = max(1, min(3, int(supersample)))
    effective = float(scale) * ss
    image = _parchment_background(gm, effective)
    draw = ImageDraw.Draw(image, "RGBA")
    _paint(draw, build_primitives(gm, edition=edition), effective,
           max(1, int(math.ceil(min_stroke * ss))) if min_stroke > 1 else 1)
    target = size or (max(1, int(gm.width * scale)), max(1, int(gm.height * scale)))
    if image.size != target:
        image = image.resize(target, Image.LANCZOS)
    return image


def _png_info(alt_text: str):
    """PNG text chunks: the picture's description, for readers and for ebooks."""
    from PIL.PngImagePlugin import PngInfo

    info = PngInfo()
    if alt_text:
        info.add_text("Description", alt_text)
    return info


def render_png(gm: GameMap, path: Path | str, scale: float = 1.0,
               supersample: int = 2, edition: str = "author") -> Path:
    """
    Rasterise a map to PNG.

    Drawn at `supersample` times the requested size and downscaled, which is
    how the lines come out smooth - Pillow has no anti-aliased line drawing.
    `edition="reader"` leaves out author-only layers. The picture's plain-
    language description (`describe_map`) is stored in the file as its
    "Description" text chunk.
    """
    scale = max(0.2, min(4.0, float(scale)))
    image = _render_image(gm, scale, supersample, edition)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(str(path), "PNG", optimize=True,
               pnginfo=_png_info(describe_map(gm, edition)))
    return path


def _draw_tracked(draw, x: float, y: float, text: str, font, colour: str,
                  tracking: float, anchor: str, halo: str = "",
                  stroke: int = 0) -> None:
    """Letter-spaced text, which Pillow cannot do natively."""
    widths = [draw.textlength(ch, font=font) for ch in text]
    total = sum(widths) + tracking * max(0, len(text) - 1)
    if anchor.startswith("r"):
        start = x - total
    elif anchor.startswith("l"):
        start = x
    else:
        start = x - total / 2.0
    vertical = "m" if anchor.endswith("m") else anchor[-1]
    # Every outline first and every letter after, or one letter's halo would
    # be painted over its neighbour.
    for outline in ((True, False) if halo and stroke else (False,)):
        cursor = start
        for ch, w in zip(text, widths):
            if outline:
                draw.text((cursor, y), ch, font=font, fill=halo,
                          anchor="l" + vertical, stroke_width=stroke,
                          stroke_fill=halo)
            else:
                draw.text((cursor, y), ch, font=font, fill=colour,
                          anchor="l" + vertical)
            cursor += w + tracking


def _dash_segments(points: Sequence[Point], dash: float,
                   gap: float) -> List[List[Point]]:
    """Split a polyline into dashes. Pillow has no dash support."""
    dash = max(1.0, dash)
    gap = max(1.0, gap)
    segments: List[List[Point]] = []
    current: List[Point] = []
    drawing = True
    remaining = dash
    for a, b in zip(points, points[1:]):
        ax, ay = a
        bx, by = b
        seg = math.hypot(bx - ax, by - ay)
        if seg < 1e-9:
            continue
        travelled = 0.0
        guard = 0
        while travelled < seg:
            guard += 1
            if guard > 4000:
                break
            step = min(remaining, seg - travelled)
            t0 = travelled / seg
            t1 = (travelled + step) / seg
            p0 = (ax + (bx - ax) * t0, ay + (by - ay) * t0)
            p1 = (ax + (bx - ax) * t1, ay + (by - ay) * t1)
            if drawing:
                if not current:
                    current.append(p0)
                current.append(p1)
            travelled += step
            remaining -= step
            if remaining <= 1e-9:
                if drawing and current:
                    segments.append(current)
                    current = []
                drawing = not drawing
                remaining = dash if drawing else gap
    if current:
        segments.append(current)
    return segments


# ==========================================================================
# SVG export
# ==========================================================================


def _svg_escape(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


def render_svg(gm: GameMap, path: Path | str, edition: str = "author") -> Path:
    """
    Write the map as SVG. `edition="reader"` leaves out author-only layers.
    The file carries the picture's plain-language description as its <title>
    and <desc>, which screen readers and ebook tools use as alt text.
    """
    palette = gm.palette()
    parts: List[str] = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{gm.width}" '
        f'height="{gm.height}" viewBox="0 0 {gm.width} {gm.height}" role="img">',
        f'<title>{_svg_escape(gm.name or "Map")}</title>',
        f'<desc>{_svg_escape(describe_map(gm, edition))}</desc>',
        f'<rect width="{gm.width}" height="{gm.height}" '
        f'fill="{palette["paper"]}"/>',
    ]

    def pts_attr(points: Sequence[Point]) -> str:
        return " ".join(f"{p[0]:.2f},{p[1]:.2f}" for p in points)

    for prim in build_primitives(gm, edition=edition):
        head = prim[0]
        if head == "polygon":
            _k, points, fill, outline, width, dash = prim
            if len(points) < 3:
                continue
            parts.append(
                f'<polygon points="{pts_attr(points)}" '
                f'fill="{fill or "none"}" stroke="{outline or "none"}" '
                f'stroke-width="{width:.2f}" stroke-linejoin="round"'
                + (' stroke-dasharray="9 6"' if dash else "") + "/>"
            )
        elif head == "line":
            _k, points, colour, width, dash = prim
            if len(points) < 2 or not colour:
                continue
            parts.append(
                f'<polyline points="{pts_attr(points)}" fill="none" '
                f'stroke="{colour}" stroke-width="{width:.2f}" '
                f'stroke-linecap="round" stroke-linejoin="round"'
                + (' stroke-dasharray="9 6"' if dash else "") + "/>"
            )
        elif head == "ellipse":
            _k, x0, y0, x1, y1, fill, outline, width = prim
            cx, cy = (x0 + x1) / 2.0, (y0 + y1) / 2.0
            rx, ry = abs(x1 - x0) / 2.0, abs(y1 - y0) / 2.0
            if rx < 0.4 or ry < 0.4:
                continue
            parts.append(
                f'<ellipse cx="{cx:.2f}" cy="{cy:.2f}" rx="{rx:.2f}" '
                f'ry="{ry:.2f}" fill="{fill or "none"}" '
                f'stroke="{outline or "none"}" stroke-width="{width:.2f}"/>'
            )
        elif head == "text":
            (x, y, text, size, colour, anchor, italic, bold, tracking,
             halo) = text_parts(prim)
            anchor_svg = {"center": "middle", "w": "start", "e": "end",
                          "n": "middle", "s": "middle"}.get(anchor, "middle")
            baseline = {"n": "hanging", "s": "auto"}.get(anchor, "central")
            parts.append(
                f'<text x="{x:.2f}" y="{y:.2f}" font-family="Georgia, serif" '
                f'font-size="{size}" fill="{colour}" '
                f'text-anchor="{anchor_svg}" dominant-baseline="{baseline}"'
                + (' font-style="italic"' if italic else "")
                + (' font-weight="bold"' if bold else "")
                + (f' letter-spacing="{tracking:.2f}"' if tracking else "")
                + (f' stroke="{halo}" stroke-width="{size * 0.26:.2f}" '
                   f'stroke-linejoin="round" paint-order="stroke"' if halo else "")
                + f">{_svg_escape(text)}</text>"
            )

    parts.append("</svg>")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(parts), encoding="utf-8")
    return path


# ==========================================================================
# Word export
# ==========================================================================


# What each kind of shape is called when a sentence lists what a map shows.
_SHAPE_NOUNS = {
    "land": "land", "water": "water", "forest": "forests",
    "mountains": "mountain ranges", "hills": "hills", "desert": "deserts",
    "swamp": "marshes", "ice": "ice and tundra", "region": "borders",
    "river": "rivers", "road": "roads", "wall": "walls", "route": "routes",
}

# How each kind of map is introduced in a description.
_MAP_PHRASE = {
    "world": "a world map", "continent": "a map of a continent",
    "region": "a regional map", "city": "a city plan",
    "building": "a building plan", "dungeon": "a dungeon map",
    "treasure": "a treasure map", "battle": "a battle plan",
}


def _list_in_words(items: Sequence[str]) -> str:
    items = list(items)
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def named_places(gm: GameMap, edition: str = "author") -> List[str]:
    """
    Every distinct name on the map, most important first, for the layers this
    edition shows: capitals and cities, then the large free-standing names
    (regions, seas), then the smaller places, then named areas.
    """
    visible = gm.visible_layers(edition)
    ranked: List[Tuple[float, float, str]] = []
    for pin in gm.pins:
        if pin.layer in visible and pin.label.strip():
            ranked.append((float(LABEL_PRIORITY.get(pin.kind, 50)),
                           -float(pin.size), pin.label.strip()))
    for label in gm.labels:
        if label.layer in visible and label.text.strip():
            ranked.append((3.5, -float(label.size), label.text.strip()))
    for shape in gm.shapes:
        if shape.layer in visible and shape.label.strip():
            ranked.append((6.5, 0.0, shape.label.strip()))
    seen: set = set()
    out: List[str] = []
    for _rank, _size, name in sorted(ranked, key=lambda r: (r[0], r[1], r[2].lower())):
        if name.lower() not in seen:
            seen.add(name.lower())
            out.append(name)
    return out


def describe_map(gm: GameMap, edition: str = "author", limit: int = 420) -> str:
    """
    A plain-language description of the picture, for alt text.

    Says what kind of map it is and how it is drawn, what it shows, how many
    places are named and which are the notable ones, and the scale. Built only
    from what `edition` shows, so the reader edition's description never names
    something on an author-only layer - and never quotes a note. Kept under
    `limit` characters, cut at a sentence.
    """
    visible = gm.visible_layers(edition)
    phrase = _MAP_PHRASE.get(gm.kind)
    if phrase is None:
        label = MAP_KINDS.get(gm.kind, gm.kind or "").lower()
        phrase = f"a {label} map" if label else "a map"
    style = STYLES.get(gm.style, {}).get("label", gm.style)
    first = phrase[0].upper() + phrase[1:]
    title = f' titled "{gm.name.strip()}"' if gm.name.strip() else ""
    sentences = [f"{first}{title}, drawn in the {style} style."]

    kinds = {s.kind for s in gm.shapes if s.layer in visible and len(s.points) >= 2}
    order = TERRAIN_ORDER + [k for k in TERRAIN if k not in TERRAIN_ORDER]
    nouns = [_SHAPE_NOUNS.get(k) or TERRAIN.get(k, {}).get("label", k).lower()
             for k in order if k in kinds]
    if nouns:
        sentences.append(f"It shows {_list_in_words(nouns)}.")

    names = named_places(gm, edition)
    if len(names) == 1:
        sentences.append(f"One place is named: {names[0]}.")
    elif names:
        sentences.append(f"{len(names)} places are named, including "
                         f"{_list_in_words(names[:6])}.")
    else:
        sentences.append("Nothing on it is named.")

    caption = gm.scale_text.strip()
    if caption:
        sentences.append(f"The scale bar reads {caption}.")

    text = ""
    for sentence in sentences:
        if len(text) + len(sentence) + 1 > limit and text:
            break
        text = f"{text} {sentence}".strip()
    return text


def render_docx(gm: GameMap, png_path: Optional[Path], out_path: Path,
                location_lookup: Optional[Dict[str, str]] = None,
                edition: str = "author") -> Path:
    """
    A Word document holding the map image plus a legend of every pin.

    Keeps maps inside the "everything is a Word document" promise, and gives
    you something printable to keep beside you while writing.

    Only what the map actually shows is listed: a hidden layer's places do not
    appear, and the reader edition also leaves out author-only layers, every
    note and the private Location-sheet column. (It used to list every pin with
    its notes and every named shape whatever the layer - a spoiler leak.) The
    picture's plain-language description is stored as its alt text.

    `png_path` is the picture to embed; give it drawn with the same `edition`,
    or pass None and this draws the right one itself.
    """
    import tempfile

    from docx.shared import Inches

    from .atomic import save_via_atomic
    from . import docxio

    visible = gm.visible_layers(edition)
    reader = edition == "reader"
    alt = describe_map(gm, edition)

    doc = docxio.sheet_document(margin=0.7)
    docxio.add_title_block(
        doc, gm.name, f"{MAP_KINDS.get(gm.kind, gm.kind)} - "
                      f"{STYLES.get(gm.style, {}).get('label', gm.style)}",
        "MAP",
    )

    with tempfile.TemporaryDirectory(prefix="nf_map_") as scratch:
        if png_path is None:
            png_path = render_png(gm, Path(scratch) / "map.png", edition=edition)
        if Path(png_path).exists():
            # 7.1in fits US Letter with 0.7in margins.
            picture = doc.add_picture(str(png_path), width=Inches(7.1))
            props = picture._inline.docPr
            props.set("descr", alt)
            props.set("title", gm.name or "Map")

        if gm.scale_text.strip():
            docxio.add_note(doc, f"Scale: {gm.scale_text}")

        pins = [p for p in gm.pins if p.layer in visible]
        if pins:
            doc.add_heading("Legend", level=2)
            rows = []
            for pin in sorted(pins, key=lambda p: (p.kind, p.label.lower())):
                row = [pin.label or "(unnamed)",
                       PIN_KINDS.get(pin.kind, pin.kind)]
                if not reader:
                    linked = ""
                    if pin.entity_id and location_lookup:
                        linked = location_lookup.get(pin.entity_id, "")
                    row += [linked, pin.notes]
                rows.append(row)
            if reader:
                docxio.add_data_table(doc, ["Place", "Type"], rows,
                                      widths=[3.6, 3.5])
            else:
                docxio.add_data_table(
                    doc, ["Place", "Type", "Location sheet", "Notes"], rows,
                    widths=[1.7, 1.2, 1.5, 2.6],
                )

        shapes = [s for s in gm.shapes if s.layer in visible]
        regions = [s for s in shapes if s.label.strip()]
        if regions:
            doc.add_heading("Named areas", level=2)
            docxio.add_data_table(
                doc, ["Name", "Kind"],
                [[s.label, TERRAIN.get(s.kind, {}).get("label", s.kind)]
                 for s in regions],
                widths=[3.0, 2.0],
            )

        if gm.notes.strip() and not reader:
            doc.add_heading("Notes", level=2)
            for para in gm.notes.split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())

        counts: Dict[str, int] = {}
        for shape in shapes:
            counts[shape.kind] = counts.get(shape.kind, 0) + 1
        if counts:
            doc.add_heading("Contents", level=2)
            summary = ", ".join(
                f"{n} x {TERRAIN.get(k, {}).get('label', k).lower()}"
                for k, n in sorted(counts.items())
            )
            labels = [l for l in gm.labels if l.layer in visible]
            doc.add_paragraph(f"{summary}. {len(pins)} pins, "
                              f"{len(labels)} labels.")

        return save_via_atomic(out_path, doc.save)


# ==========================================================================
# Persistence
# ==========================================================================

MAP_SUFFIX = ".map.json"


def map_path(folder: Path, gm: GameMap) -> Path:
    return Path(folder) / f"{safe_filename(gm.name, 'Map')}{MAP_SUFFIX}"


def save_map(folder: Path, gm: GameMap) -> Path:
    gm.touch()
    path = map_path(folder, gm)
    write_json_atomic(path, gm.to_json())
    return path


def load_map(path: Path) -> Optional[GameMap]:
    data = read_json(path)
    if not isinstance(data, dict):
        return None
    try:
        return GameMap.from_json(data)
    except (TypeError, ValueError):
        return None


def list_maps(folder: Path) -> List[Tuple[str, Path]]:
    folder = Path(folder)
    if not folder.is_dir():
        return []
    out: List[Tuple[str, Path]] = []
    for path in sorted(folder.glob(f"*{MAP_SUFFIX}")):
        data = read_json(path, {}) or {}
        out.append((data.get("name") or path.name[:-len(MAP_SUFFIX)], path))
    return out


def delete_map(path: Path) -> bool:
    try:
        Path(path).unlink(missing_ok=True)
        return True
    except OSError:
        return False


# ==========================================================================
# Starter content
# ==========================================================================


def starter_map(name: str = "The Known World", kind: str = "world",
                style: str = "parchment") -> GameMap:
    """
    A new map with a plausible coastline already on it.

    An empty canvas is paralysing; something to push around is not. Generated
    from a fixed seed so it is the same every time and can be deleted in one
    click if unwanted.
    """
    gm = GameMap(name=name, kind=kind, style=style)
    if kind in ("city", "building", "dungeon", "battle"):
        gm.grid = "square"
        gm.grid_size = 60
        gm.compass = kind != "building"
        return gm

    rng = random.Random(11)
    cx, cy = gm.width * 0.47, gm.height * 0.52
    rx, ry = gm.width * 0.31, gm.height * 0.33
    coast: List[Point] = []
    steps = 46
    for i in range(steps):
        a = 2 * math.pi * i / steps
        wobble = 1.0 + 0.22 * math.sin(a * 3.0 + 0.7) + rng.uniform(-0.09, 0.09)
        coast.append((cx + math.cos(a) * rx * wobble,
                      cy + math.sin(a) * ry * wobble))
    gm.shapes.append(Shape(kind="land", points=coast, closed=True,
                           label="", layer="Base"))

    gm.shapes.append(Shape(
        kind="mountains", closed=False, layer="Base",
        points=[(cx - rx * 0.45, cy - ry * 0.42), (cx - rx * 0.1, cy - ry * 0.52),
                (cx + rx * 0.25, cy - ry * 0.38), (cx + rx * 0.5, cy - ry * 0.16)],
    ))
    gm.shapes.append(Shape(
        kind="forest", closed=True, layer="Base",
        points=[(cx - rx * 0.62, cy + ry * 0.05), (cx - rx * 0.2, cy - ry * 0.05),
                (cx - rx * 0.08, cy + ry * 0.38), (cx - rx * 0.55, cy + ry * 0.46)],
    ))
    gm.shapes.append(Shape(
        kind="river", closed=False, layer="Base",
        points=[(cx - rx * 0.05, cy - ry * 0.44), (cx + rx * 0.04, cy - ry * 0.1),
                (cx - rx * 0.06, cy + ry * 0.22), (cx + rx * 0.12, cy + ry * 0.62),
                (cx + rx * 0.3, cy + ry * 0.95)],
    ))
    gm.pins.append(Pin(x=cx + rx * 0.14, y=cy + ry * 0.6, kind="capital",
                       label="Capital", label_side="e", size=8))
    gm.pins.append(Pin(x=cx - rx * 0.5, y=cy + ry * 0.3, kind="town",
                       label="A town", label_side="w"))
    gm.pins.append(Pin(x=cx + rx * 0.55, y=cy - ry * 0.05, kind="ruin",
                       label="Ruins", label_side="e"))
    gm.labels.append(MapLabel(x=cx, y=cy + ry * 1.28, text="THE OCEAN",
                              size=20, italic=True, tracking=6.0))
    gm.scale_text = "100 leagues"
    return gm


NAME_SYLLABLES = {
    "northern": (["Bran", "Thor", "Sker", "Vald", "Hald", "Grim", "Fjor",
                  "Ulf", "Stein", "Orm", "Rag", "Sig"],
                 ["dal", "vik", "holm", "gard", "fell", "borg", "mark",
                  "stad", "ness", "fjord", "heim", "by"]),
    "southern": (["Cala", "Meri", "Sera", "Vala", "Alma", "Tara", "Sola",
                  "Bela", "Cora", "Lira", "Nera", "Mira"],
                 ["nova", "mar", "vera", "londe", "riva", "sette", "monte",
                  "corte", "valle", "porto", "bella", "sana"]),
    "elvish": (["Ael", "Cel", "Ely", "Fin", "Gal", "Ith", "Lor", "Mith",
                "Nim", "Sil", "Thal", "Vae"],
               ["andor", "ariel", "wen", "rond", "aloth", "ithil", "dor",
                "las", "riel", "orn", "eth", "amar"]),
    "harsh": (["Kra", "Zor", "Gnash", "Vrak", "Mog", "Skul", "Thrag",
               "Ur", "Bok", "Drez", "Hak", "Nur"],
              ["gul", "zar", "dun", "rok", "grim", "mor", "kath", "thul",
               "gash", "var", "nak", "dread"]),
    "plain": (["Ash", "Black", "Cold", "Green", "Grey", "High", "Long",
               "North", "Oak", "Salt", "Stone", "White", "Red", "Deep"],
              ["ford", "bridge", "field", "wood", "hill", "water", "gate",
               "haven", "moor", "brook", "ridge", "reach", "hollow", "cross"]),
}


#: The shipped lists, kept so the editable file can always be restored.
DEFAULT_NAME_SYLLABLES = {
    key: (list(start), list(end))
    for key, (start, end) in NAME_SYLLABLES.items()
}


#: Which style names which kind of thing. A kingdom rarely sounds like a
#: village, and a river almost never sounds like either - so each gets its own
#: entry rather than everything sharing one list.
NAME_ROLES: Dict[str, str] = {
    "settlement": "plain",
    "realm": "northern",
    "region": "northern",
    "water": "elvish",
    "river": "elvish",
}

DEFAULT_NAME_ROLES = dict(NAME_ROLES)

ROLE_LABELS = {
    "settlement": "Towns, cities and keeps",
    "realm": "Kingdoms and realms",
    "region": "Regions and provinces",
    "water": "Seas and oceans",
    "river": "Rivers and lakes",
}


def name_styles_file(folder: Path | str) -> Path:
    return Path(folder) / "Name Styles.json"


def name_for(role: str, fallback: str = "", seed: Optional[int] = None) -> str:
    """
    A name for one kind of thing, in whatever style that kind uses.

    `fallback` is the map's own style, used when the role has no entry - so a
    writer who never touches the roles still gets a consistent world. Give a
    `seed` and the same seed always gives the same name.
    """
    style = NAME_ROLES.get(role) or fallback
    if style not in NAME_SYLLABLES:
        style = fallback if fallback in NAME_SYLLABLES else ""
    return generate_name(style or next(iter(NAME_SYLLABLES), "plain"), seed)


def load_name_styles(folder: Path | str) -> Path:
    """
    Read the writer's own name lists, creating the starter file if absent.

    The generator used to be five hardcoded lists, which is fine until your
    world does not sound like any of them. This is a plain JSON file in the
    Maps folder: add a style, delete one, replace every syllable with your own.
    Anything malformed is ignored in favour of the shipped lists rather than
    breaking name generation.
    """
    path = name_styles_file(folder)
    if not path.exists():
        try:
            write_json_atomic(path, {
                "_note": "Each style has a list of beginnings and a list of "
                         "endings. A name is one of each, joined. Add your own "
                         "styles, or replace these entirely.",
                "_roles_note": "Which style names which kind of thing. Change "
                               "a value to any style name below - or to the "
                               "same one everywhere, if your world sounds "
                               "consistent.",
                "_roles": dict(DEFAULT_NAME_ROLES),
                **{key: {"start": start, "end": end}
                   for key, (start, end) in DEFAULT_NAME_SYLLABLES.items()},
            })
        except OSError:
            return path

    data = read_json(path, None)
    if not isinstance(data, dict):
        return path
    merged: Dict[str, Tuple[List[str], List[str]]] = {}
    for key, value in data.items():
        if key.startswith("_") or not isinstance(value, dict):
            continue
        start = value.get("start")
        end = value.get("end")
        if isinstance(start, list) and isinstance(end, list) and start and end:
            merged[str(key)] = ([str(s) for s in start if str(s).strip()],
                                [str(e) for e in end if str(e).strip()])
    if merged:
        NAME_SYLLABLES.clear()
        NAME_SYLLABLES.update(merged)

    roles = data.get("_roles")
    NAME_ROLES.clear()
    NAME_ROLES.update(DEFAULT_NAME_ROLES)
    if isinstance(roles, dict):
        for role, style in roles.items():
            if isinstance(style, str) and str(role) in DEFAULT_NAME_ROLES:
                NAME_ROLES[str(role)] = style
    # A role pointing at a style the writer deleted would silently fall back
    # to the map's own style; better to point it somewhere real.
    available = sorted(NAME_SYLLABLES)
    for role, style in list(NAME_ROLES.items()):
        if style not in NAME_SYLLABLES:
            NAME_ROLES[role] = available[0] if available else "plain"
    return path


def name_styles() -> List[str]:
    return sorted(NAME_SYLLABLES)


def generate_name(flavour: str = "plain", seed: Optional[int] = None) -> str:
    """A place-name generator. Trivial, and saves a surprising amount of time."""
    if flavour not in NAME_SYLLABLES and NAME_SYLLABLES:
        flavour = next(iter(NAME_SYLLABLES))
    first, second = NAME_SYLLABLES.get(
        flavour, next(iter(NAME_SYLLABLES.values())))
    rng = random.Random(seed)
    name = rng.choice(first) + rng.choice(second)
    if flavour == "plain" and rng.random() < 0.25 and "plain" in NAME_SYLLABLES:
        name += " " + rng.choice(["Keep", "Cross", "End", "Mill", "Watch"])
    return name


def generate_names(flavour: str = "plain", count: int = 12) -> List[str]:
    out: List[str] = []
    guard = 0
    while len(out) < count and guard < count * 20:
        guard += 1
        candidate = generate_name(flavour)
        if candidate not in out:
            out.append(candidate)
    return out
