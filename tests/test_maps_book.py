"""
Book-ready maps: the scale that tells the truth, secrets that stay secret, print
and ebook export, the legend, and the links from a pin to the map inside it.

No window is opened here (one class near the bottom opens a hidden Tk canvas,
and skips itself where there is no display), so this is fast. Every check was
proven able to fail by breaking the thing it guards; the docstrings say what.
"""

import math
import tempfile
import unittest
import zlib
from pathlib import Path

from tests import require_isolation

require_isolation()

from novelforge import mapgen, mapmaker as mm, mapstory  # noqa: E402


def blank(width=1600, **kw):
    """A map with nothing on it but its furniture."""
    kw.setdefault("height", round(width * 0.6875))
    return mm.GameMap(name="Book", width=width, layers=[mm.Layer(name="Base")], **kw)


def drawn_bar_px(gm):
    """
    How long the scale bar is in the picture, read off the display list.

    Deliberately does not ask `scale_bar` (that would be checking a function
    against itself): it finds the bar's four blocks by where the design puts
    them - left margin 4.5% of the width, 5.2% of the height up from the bottom.
    """
    left = gm.width * 0.045
    top = gm.height - gm.height * 0.052
    blocks = [p for p in mm.build_primitives(gm)
              if p[0] == "polygon" and len(p[1]) == 4
              and abs(p[1][0][1] - top) < 1e-6 and p[1][0][0] >= left - 1e-6]
    assert len(blocks) == 4, f"expected the scale bar's 4 blocks, found {len(blocks)}"
    return max(pt[0] for p in blocks for pt in p[1]) - min(
        pt[0] for p in blocks for pt in p[1])


# The number in each caption, typed here by hand rather than parsed, so the
# test does not lean on the parser it is checking.
CAPTIONS = (("200 leagues", 200.0), ("100 miles", 100.0), ("50 km", 50.0),
            ("1,000 feet", 1000.0), ("3 AU", 3.0), ("12.5 miles", 12.5),
            ("5 leagues (15 miles)", 5.0))


class ScaleTellsTheTruth(unittest.TestCase):
    """The bar is drawn 16% of the width (at most 230 px); mapstory assumed 20%."""

    def test_a_distance_is_what_the_drawn_bar_says_it_is(self):
        # Breaking it: putting `width * 0.2` back into mapstory.read_scale
        # fails every width here (the drawn bar is 192, 230, 230 and 230 px
        # long at these widths; the old guess was 240, 320, 360 and 480).
        for width in (1200, 1600, 1800, 2400):
            for caption, amount in CAPTIONS:
                gm = blank(width, scale_text=caption)
                bar = drawn_bar_px(gm)
                for apart in (bar, bar * 0.5, 137.0, 611.0):
                    a = mm.Pin(x=100, y=200)
                    b = mm.Pin(x=100 + apart * 0.6, y=200 + apart * 0.8)  # 3-4-5
                    self.assertAlmostEqual(
                        mapstory.distance(gm, a, b), apart / bar * amount,
                        places=6, msg=f"{width}px wide, '{caption}', {apart:.0f}px apart")

    def test_two_pins_one_bar_apart_are_exactly_the_caption_apart(self):
        for width in (1200, 1600, 1800, 2400):
            gm = blank(width, scale_text="200 leagues")
            bar = drawn_bar_px(gm)
            a, b = mm.Pin(x=50, y=50), mm.Pin(x=50 + bar, y=50)
            self.assertAlmostEqual(mapstory.distance(gm, a, b), 200.0, places=6)

    def test_the_old_guess_of_a_fifth_of_the_width_is_gone(self):
        gm = blank(1800, scale_text="200 leagues")
        scale = mapstory.read_scale(gm)
        self.assertAlmostEqual(scale.bar_pixels, drawn_bar_px(gm), places=6)
        self.assertNotAlmostEqual(scale.bar_pixels, 1800 * 0.2, places=0)
        self.assertEqual(round(drawn_bar_px(gm)), 230)        # capped, as drawn since 2.2

    def test_old_maps_are_drawn_exactly_as_before(self):
        # The default bar is 16% of the width, capped at 230 px - what every map
        # saved before this change was drawn with, so nothing already made moves.
        for width, expected in ((1200, 192.0), (1400, 224.0), (1600, 230.0),
                                (1800, 230.0), (2400, 230.0)):
            gm = blank(width, scale_text="100 leagues")
            self.assertAlmostEqual(drawn_bar_px(gm), expected, places=6)

    def test_names_keep_off_the_bar_that_is_actually_drawn(self):
        # layout_pin_labels reserved `min(width * 0.16, 230)` itself, separately
        # from the drawing. A pin whose name sits just past the drawn bar must
        # keep its preferred side; if the reservation grows to a fifth of the
        # width (360 px) that side is blocked and the name is moved or dropped.
        gm = blank(1800, scale_text="200 leagues")
        left = gm.width * 0.045
        top = gm.height - gm.height * 0.052
        pin = mm.Pin(x=left + 300, y=top - 5, kind="town", label="Ash",
                     label_side="e")
        gm.pins.append(pin)
        self.assertEqual(mm.layout_pin_labels(gm)[pin.id], ("e", True))

    def test_a_bar_with_its_own_length_is_read_back_the_same_way(self):
        gm = blank(1800, scale_text="80 miles", scale_px=120.0)
        self.assertAlmostEqual(drawn_bar_px(gm), 120.0, places=6)
        self.assertAlmostEqual(mapstory.read_scale(gm).units_per_pixel, 80 / 120.0)

    def test_no_caption_no_bar_and_no_invented_distance(self):
        gm = blank(1600)
        self.assertFalse(mm.scale_bar(gm).label)
        self.assertFalse(mapstory.read_scale(gm).known)
        self.assertIsNone(mapstory.distance(gm, mm.Pin(), mm.Pin(x=90, y=0)))
        words = blank(1600, scale_text="a long walk")
        self.assertTrue(mm.scale_bar(words).label)             # still drawn ...
        self.assertFalse(mapstory.read_scale(words).known)     # ... but no number

    def test_a_map_from_an_older_file_still_opens_and_measures(self):
        old = {"id": "map_old", "name": "Old", "width": 1600, "height": 1100,
               "scale_text": "100 leagues", "layers": [{"name": "Base"}],
               "pins": [{"id": "p1", "x": 10, "y": 10, "label": "A"}]}
        gm = mm.GameMap.from_json(old)
        self.assertEqual(gm.scale_px, 0.0)
        self.assertFalse(gm.legend)
        self.assertEqual(gm.parent, {})
        self.assertAlmostEqual(mapstory.read_scale(gm).bar_pixels, 230.0)
        back = mm.GameMap.from_json(gm.to_json())
        self.assertEqual(back.scale_px, 0.0)


class NiceNumbers(unittest.TestCase):
    def test_floor_ceil_and_round_are_ones_twos_and_fives(self):
        floor = [(19.2, 10), (20, 20), (49.9, 20), (50, 50), (99, 50),
                 (100, 100), (0.037, 0.02), (0.5, 0.5), (1, 1), (7.4e3, 5000)]
        for value, expected in floor:
            self.assertAlmostEqual(mm.nice_number(value, "floor"), expected,
                                   msg=f"floor({value})", places=9)
        for value, expected in ((21, 50), (50, 50), (51, 100), (0.11, 0.2)):
            self.assertAlmostEqual(mm.nice_number(value, "ceil"), expected,
                                   msg=f"ceil({value})", places=9)
        for value, expected in ((3, 2), (4, 5), (140, 100), (160, 200)):
            self.assertAlmostEqual(mm.nice_number(value, "round"), expected,
                                   msg=f"round({value})", places=9)
        self.assertEqual(mm.nice_number(0), 0.0)
        self.assertEqual(mm.nice_number(-4), 0.0)

    def test_every_answer_is_a_one_two_or_five(self):
        for i in range(1, 400):
            value = 1.037 ** i
            for mode in ("floor", "ceil", "round"):
                answer = mm.nice_number(value, mode)
                lead = answer / 10.0 ** math.floor(math.log10(answer) + 1e-9)
                self.assertTrue(any(abs(lead - s) < 1e-6 for s in (1, 2, 5)),
                                f"{mode}({value}) = {answer}")


class UnitsAndCaptions(unittest.TestCase):
    MILES = (("10 ft", 10 / 5280), ("10 feet", 10 / 5280), ("3 yards", 9 / 5280),
             ("2 yd", 6 / 5280), ("40 paces", 100 / 5280), ("50 metres", 0.0310686),
             ("50 m", 0.0310686), ("2 km", 1.242742), ("2 kilometers", 1.242742),
             ("1 mile", 1.0), ("4 mi", 4.0), ("2 leagues", 6.0),
             ("3 days", 60.0))

    def test_every_unit_reads_and_converts(self):
        for caption, miles in self.MILES:
            gm = blank(1600, scale_text=caption)
            scale = mapstory.read_scale(gm)
            self.assertTrue(scale.known, caption)
            amount = float(caption.split()[0])
            self.assertAlmostEqual(scale.to_miles(amount), miles, delta=miles * 1e-3,
                                   msg=caption)

    def test_the_stars_are_measured_in_their_own_units(self):
        one = {"4 parsecs": 4, "4 pc": 4, "3 AU": 3, "3 au": 3,
               "6 light years": 6, "6 light-years": 6, "6 ly": 6}
        for caption, amount in one.items():
            scale = mapstory.read_scale(blank(1600, scale_text=caption))
            self.assertTrue(scale.known, caption)
            self.assertTrue(scale.astronomical, caption)
            self.assertIsNotNone(scale.to_miles(amount), caption)
        # a parsec is about 19 trillion miles - nobody wants that printed
        gm = blank(1600, scale_text="4 parsecs")
        text = mapstory.describe_distance(gm, mm.Pin(x=0, y=0), mm.Pin(x=230, y=0))
        self.assertEqual(text, "4 parsecs")

    def test_a_caption_with_thousands_and_decimals_and_trailing_words(self):
        self.assertEqual(mm.parse_scale_caption("1,500 miles"), (1500.0, "miles"))
        self.assertEqual(mm.parse_scale_caption("12.5 km"), (12.5, "km"))
        self.assertEqual(mm.parse_scale_caption("200 leagues wide"),
                         (200.0, "leagues"))
        self.assertEqual(mm.parse_scale_caption("3 Light Years"),
                         (3.0, "light years"))
        self.assertEqual(mm.parse_scale_caption("75"), (75.0, "units"))
        self.assertIsNone(mm.parse_scale_caption("a long walk"))
        self.assertIsNone(mm.parse_scale_caption("0 miles"))
        self.assertIsNone(mm.parse_scale_caption(""))

    def test_captions_agree_with_their_number(self):
        self.assertEqual(mm.scale_caption(1, "miles"), "1 mile")
        self.assertEqual(mm.scale_caption(5, "mile"), "5 miles")
        self.assertEqual(mm.scale_caption(10, "ft"), "10 ft")
        self.assertEqual(mm.scale_caption(1, "parsecs"), "1 parsec")
        self.assertEqual(mm.scale_caption(1000, "leagues"), "1,000 leagues")
        self.assertEqual(mm.scale_caption(0.5, "miles"), "0.5 miles")
        self.assertEqual(mm.scale_caption(2, "AU"), "2 AU")
        self.assertEqual(mm.scale_caption(3, "quarrels"), "3 quarrels")

    def test_a_scale_that_is_only_a_number_is_still_a_scale(self):
        scale = mapstory.read_scale(blank(1600, scale_text="75"))
        self.assertTrue(scale.known)
        self.assertEqual(scale.unit, "units")
        self.assertIsNone(scale.to_miles(75))                 # no unit, no miles


class ApplyScale(unittest.TestCase):
    CASES = (("ft", 5, 60, 1600), ("ft", 5, 40, 1200), ("miles", 10, 80, 1800),
             ("miles", 3, 50, 2400), ("parsecs", 1, 90, 1600),
             ("leagues", 25, 120, 1600), ("km", 0.5, 30, 1200))

    def test_the_caption_the_bar_and_the_grid_cannot_disagree(self):
        # Breaking it: leave `scale_px` unset in apply_scale and the bar is drawn
        # at its default length while the grid square still means per_cell.
        for unit, per_cell, cell, width in self.CASES:
            gm = blank(width, grid="square", grid_size=cell)
            bar = mm.apply_scale(gm, unit, per_cell)
            what = f"{per_cell} {unit} per {cell}px on {width}px"
            self.assertTrue(bar.known, what)
            # the caption carries a round number ...
            lead = bar.value / 10.0 ** math.floor(math.log10(bar.value) + 1e-9)
            self.assertTrue(any(abs(lead - s) < 1e-6 for s in (1, 2, 5)), what)
            # ... the picture draws exactly the bar that number needs ...
            self.assertAlmostEqual(drawn_bar_px(gm), bar.value / per_cell * cell,
                                   places=6, msg=what)
            # ... it fits where a bar goes, without shrinking to a stub ...
            room = min(width * 0.16, 230.0)
            self.assertLessEqual(drawn_bar_px(gm), room + 1e-6, what)
            self.assertGreaterEqual(drawn_bar_px(gm), room * 0.4 - 1e-6, what)
            # ... and one grid square is worth `per_cell`, as read by the story.
            a, b = mm.Pin(x=0, y=0), mm.Pin(x=cell, y=0)
            self.assertAlmostEqual(mapstory.distance(gm, a, b), per_cell,
                                   places=6, msg=what)
            self.assertAlmostEqual(mm.cell_units(gm), per_cell, places=6, msg=what)

    def test_five_feet_a_square_on_a_dungeon(self):
        gm = blank(1600, grid="square", grid_size=60)
        mm.apply_scale(gm, "ft", 5)
        self.assertEqual(gm.scale_text, "10 ft")
        self.assertAlmostEqual(gm.scale_px, 120.0)
        self.assertEqual(gm.grid_size, 60)

    def test_giving_the_square_size_sets_the_grid_too(self):
        gm = blank(1600, grid="square", grid_size=80)
        mm.apply_scale(gm, "ft", 5, cell_px=57.6)
        self.assertEqual(gm.grid_size, 58)                    # whole pixels
        a, b = mm.Pin(x=0, y=0), mm.Pin(x=58, y=0)
        self.assertAlmostEqual(mapstory.distance(gm, a, b), 5.0, places=6)

    def test_a_scale_is_saved_and_comes_back(self):
        gm = blank(1600, grid="square", grid_size=60)
        mm.apply_scale(gm, "ft", 5)
        again = mm.GameMap.from_json(gm.to_json())
        self.assertEqual((again.scale_text, again.scale_px, again.grid_size),
                         (gm.scale_text, gm.scale_px, gm.grid_size))
        self.assertAlmostEqual(drawn_bar_px(again), drawn_bar_px(gm))

    def test_nonsense_is_refused(self):
        gm = blank(1600)
        for bad in (0, -3, float("nan"), float("inf")):
            with self.assertRaises(ValueError):
                mm.apply_scale(gm, "miles", bad)

    def test_mapstory_offers_the_same_helper(self):
        self.assertIs(mapstory.apply_scale, mm.apply_scale)


class OtherScaleUsers(unittest.TestCase):
    def test_a_generated_world_measures_in_leagues_and_says_so(self):
        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 7))
        bar = drawn_bar_px(gm)
        text = mapstory.describe_distance(gm, mm.Pin(x=0, y=0),
                                          mm.Pin(x=bar, y=0))
        self.assertEqual(text, "200 leagues (about 600 miles)")


def square(x, y, size):
    return [(x, y), (x + size, y), (x + size, y + size), (x, y + size)]


def secret_map():
    """
    A small world with a secret layer (author-only), a hidden layer and notes.

    Base: Harrowgate (capital), Dunmere, Whisperwood (a named forest), "The Grey
    Sea". Spoilers (author-only): the Vault of Kings, Lost Isle, "The Void".
    Hidden (switched off): Old Fort and Buried Ruins.
    """
    gm = mm.GameMap(
        name="The Broken Coast", kind="world", style="parchment", width=800,
        height=550, scale_text="50 miles",
        notes="The king is a fraud.\n\nThe vault holds the real crown.",
        layers=[mm.Layer(name="Base"),
                mm.Layer(name="Spoilers", author_only=True),
                mm.Layer(name="Hidden", visible=False)])
    gm.shapes = [
        mm.Shape(kind="land", layer="Base", points=square(60, 90, 460)),
        mm.Shape(kind="forest", layer="Base", label="Whisperwood",
                 points=square(150, 200, 160)),
        mm.Shape(kind="land", layer="Spoilers", label="Lost Isle",
                 points=square(560, 300, 190)),
        mm.Shape(kind="land", layer="Hidden", label="Buried Ruins",
                 points=square(560, 60, 90)),
    ]
    gm.pins = [
        mm.Pin(x=200, y=170, kind="capital", label="Harrowgate", layer="Base",
               notes="A tunnel runs under the well.", entity_id="loc_1"),
        mm.Pin(x=420, y=330, kind="town", label="Dunmere", layer="Base"),
        mm.Pin(x=650, y=390, kind="treasure", label="Vault of Kings",
               layer="Spoilers", notes="The crown is here."),
        mm.Pin(x=600, y=90, kind="ruin", label="Old Fort", layer="Hidden",
               notes="Forgotten garrison."),
    ]
    gm.labels = [
        mm.MapLabel(x=400, y=500, text="The Grey Sea", size=22, layer="Base"),
        mm.MapLabel(x=700, y=500, text="The Void", size=18, layer="Spoilers"),
    ]
    # Ids seed the trees, and are random unless set: two calls must draw alike.
    for i, item in enumerate(gm.shapes + gm.pins + gm.labels):
        item.id = f"item_{i}"
    return gm


def texts(prims):
    return [p[3] for p in prims if p[0] == "text"]


class SecretLayers(unittest.TestCase):
    def test_a_layer_is_public_unless_marked_author_only(self):
        self.assertFalse(mm.Layer().author_only)
        gm = secret_map()
        again = mm.GameMap.from_json(gm.to_json())
        self.assertEqual([l.author_only for l in again.layers], [False, True, False])
        old = mm.GameMap.from_json({"layers": [{"name": "Base", "visible": True}]})
        self.assertFalse(old.layers[0].author_only)

    def test_the_editions_see_different_layers(self):
        gm = secret_map()
        self.assertEqual(gm.visible_layers(), {"Base", "Spoilers"})
        self.assertEqual(gm.visible_layers("author"), {"Base", "Spoilers"})
        self.assertEqual(gm.visible_layers("reader"), {"Base"})
        with self.assertRaises(ValueError):
            gm.visible_layers("Reader")           # a typo must not mean "author"

    def test_the_reader_display_list_has_no_secrets(self):
        # Breaking it: make visible_layers ignore author_only, and the reader's
        # list names the vault.
        gm = secret_map()
        author = texts(mm.build_primitives(gm))
        reader = texts(mm.build_primitives(gm, edition="reader"))
        for name in ("Vault of Kings", "The Void", "Lost Isle"):
            self.assertIn(name, author)
            self.assertNotIn(name, reader)
        for name in ("Harrowgate", "Dunmere", "The Grey Sea", "Whisperwood"):
            self.assertIn(name, reader)
        for name in ("Old Fort", "Buried Ruins"):        # switched off: in neither
            self.assertNotIn(name, author)
            self.assertNotIn(name, reader)

    def test_asking_for_one_edition_never_serves_the_other_from_the_cache(self):
        gm = secret_map()
        first = mm.build_primitives(gm)
        reader = mm.build_primitives(gm, edition="reader")
        again = mm.build_primitives(gm)
        self.assertEqual(first, again)
        self.assertNotEqual(first, reader)
        self.assertEqual(mm.build_primitives(gm, edition="reader"), reader)

    def test_a_default_call_is_still_the_authors_copy(self):
        gm = secret_map()
        self.assertEqual(mm.build_primitives(gm),
                         mm.build_primitives(gm, edition="author"))
        gm2 = mapgen.generate(mapgen.preset("Classic fantasy world", 3))
        self.assertEqual(mm.build_primitives(gm2),
                         mm.build_primitives(gm2, edition="author"))

    def test_a_map_with_no_secrets_reads_the_same_in_both_editions(self):
        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 3))
        self.assertEqual(mm.build_primitives(gm),
                         mm.build_primitives(gm, edition="reader"))


class SecretsStayOutOfTheExports(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.folder.name)

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def test_the_reader_svg_leaves_them_out_and_describes_only_what_it_shows(self):
        gm = secret_map()
        author = mm.render_svg(gm, self.dir / "a.svg").read_text(encoding="utf-8")
        reader = mm.render_svg(gm, self.dir / "r.svg", edition="reader"
                               ).read_text(encoding="utf-8")
        self.assertIn("Vault of Kings", author)
        self.assertNotIn("Vault of Kings", reader)
        self.assertNotIn("Lost Isle", reader)
        self.assertIn("<desc>", reader)
        self.assertIn("Harrowgate", reader)

    def test_the_reader_picture_is_the_picture_without_that_layer(self):
        from PIL import Image

        gm = secret_map()
        for layer_edition in ("author", "reader"):
            mm.render_png(gm, self.dir / f"{layer_edition}.png", scale=0.5,
                          edition=layer_edition)
        without = secret_map()
        without.layers[1].visible = False                # the same, layer off
        mm.render_png(without, self.dir / "off.png", scale=0.5)
        with Image.open(self.dir / "reader.png") as r, \
                Image.open(self.dir / "off.png") as off, \
                Image.open(self.dir / "author.png") as a:
            self.assertEqual(r.convert("RGB").tobytes(), off.convert("RGB").tobytes())
            self.assertNotEqual(r.convert("RGB").tobytes(), a.convert("RGB").tobytes())

    def word_text(self, path):
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts += [cell.text for cell in row.cells]
        return "\n".join(parts), doc

    def test_the_word_file_lists_only_what_is_on_the_map(self):
        # The leak: render_docx listed every pin (notes and all) and every named
        # shape whatever its layer. Breaking it: list `gm.pins` again.
        gm = secret_map()
        text, _ = self.word_text(mm.render_docx(
            gm, None, self.dir / "author.docx", {"loc_1": "Harrowgate Keep"}))
        for shown in ("Harrowgate", "A tunnel runs under the well.", "Vault of Kings",
                      "The crown is here.", "Harrowgate Keep", "Whisperwood",
                      "The king is a fraud.", "Lost Isle"):
            self.assertIn(shown, text)
        for hidden in ("Old Fort", "Forgotten garrison.", "Buried Ruins"):
            self.assertNotIn(hidden, text, "a switched-off layer leaked")

    def test_the_reader_word_file_has_no_secrets_and_no_notes(self):
        gm = secret_map()
        text, _ = self.word_text(mm.render_docx(
            gm, None, self.dir / "reader.docx", {"loc_1": "Harrowgate Keep"},
            edition="reader"))
        for shown in ("Harrowgate", "Dunmere", "Whisperwood", "50 miles"):
            self.assertIn(shown, text)
        for hidden in ("Vault of Kings", "The crown is here.", "Lost Isle",
                       "Old Fort", "Buried Ruins", "A tunnel runs under the well.",
                       "The king is a fraud.", "Harrowgate Keep", "Location sheet",
                       "Notes"):
            self.assertNotIn(hidden, text, hidden)
        self.assertIn("2 pins, 1 labels", text)          # counts what is shown

    def test_the_picture_in_the_word_file_carries_alt_text(self):
        gm = secret_map()
        for edition in ("author", "reader"):
            _, doc = self.word_text(mm.render_docx(
                gm, None, self.dir / f"alt_{edition}.docx", edition=edition))
            self.assertEqual(len(doc.inline_shapes), 1)
            props = doc.inline_shapes[0]._inline.docPr
            self.assertEqual(props.get("descr"), mm.describe_map(gm, edition))
            self.assertTrue(props.get("descr"))
        reader = doc.inline_shapes[0]._inline.docPr.get("descr")
        self.assertNotIn("Vault", reader)

    def test_a_png_carries_its_description(self):
        from PIL import Image

        gm = secret_map()
        mm.render_png(gm, self.dir / "d.png", scale=0.3, edition="reader")
        with Image.open(self.dir / "d.png") as image:
            image.load()
            self.assertEqual(image.text.get("Description"),
                             mm.describe_map(gm, "reader"))


class AltText(unittest.TestCase):
    def test_it_says_what_kind_of_map_what_it_shows_and_who_lives_there(self):
        # Breaking it: build the description from every layer, not the
        # edition's, and the reader's alt text names the vault.
        gm = secret_map()
        author, reader = mm.describe_map(gm), mm.describe_map(gm, "reader")
        self.assertTrue(author.startswith('A world map titled "The Broken Coast"'))
        self.assertIn("Parchment", author)
        self.assertIn("land and forests", author)
        self.assertIn("7 places are named, including Harrowgate", author)
        self.assertIn("4 places are named, including Harrowgate", reader)
        self.assertNotIn("Vault of Kings", reader)
        self.assertNotIn("Lost Isle", reader)
        self.assertIn("The scale bar reads 50 miles.", reader)
        self.assertNotIn("fraud", author + reader)       # never a note
        self.assertNotIn("tunnel", author + reader)

    def test_capitals_lead_and_the_list_is_short(self):
        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 3))
        text = mm.describe_map(gm)
        capital = next(p.label for p in gm.pins if p.kind == "capital")
        self.assertIn(f"including {capital}", text)
        self.assertLessEqual(len(text), 420)

    def test_an_empty_or_odd_map_still_gets_a_sentence(self):
        text = mm.describe_map(mm.GameMap(name="", kind=""))
        self.assertTrue(text.startswith("A map, drawn in"), text)
        self.assertIn("Nothing on it is named.", text)
        self.assertTrue(mm.describe_map(mm.GameMap(name="x", kind="starfield"))
                        .startswith('A starfield map titled "x"'))
        one = mm.GameMap(name="Solo", kind="dungeon")
        one.pins.append(mm.Pin(label="The Pit", kind="danger"))
        self.assertIn("One place is named: The Pit.", mm.describe_map(one))
        self.assertTrue(mm.describe_map(one).startswith("A dungeon map"))


def hex_rgb(colour):
    return tuple(int(colour[i:i + 2], 16) for i in (1, 3, 5))


def luma(colour):
    """Grey level 0..1 the way Pillow's convert("L") sees a colour."""
    r, g, b = hex_rgb(colour)
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255.0


class PrintStyle(unittest.TestCase):
    """STYLES["print"]: black ink on white, and greys a black-and-white press keeps apart."""

    ADJACENT = (
        ("sea", "land"),
        ("land", "water"), ("land", "forest"), ("land", "desert"),
        ("land", "swamp"), ("land", "ice"),
        ("forest", "swamp"), ("forest", "desert"), ("forest", "ice"),
        ("desert", "swamp"), ("desert", "ice"), ("swamp", "ice"),
        ("water", "forest"), ("water", "desert"), ("water", "swamp"),
        ("water", "ice"), ("sea", "ice"), ("sea", "forest"), ("sea", "desert"),
        ("sea", "swamp"),
    )

    def fills(self):
        style = mm.STYLES["print"]
        out = dict(style["terrain"])
        out["sea"] = style["paper"]
        return out

    def test_it_is_a_style_the_editor_can_offer(self):
        self.assertIn("print", mm.STYLES)
        self.assertTrue(mm.STYLES["print"]["label"])
        self.assertEqual(mm.GameMap(style="print").palette(), mm.STYLES["print"])

    def test_every_colour_in_it_is_a_grey(self):
        def colours(value):
            if isinstance(value, str):
                yield value
            elif isinstance(value, dict):
                for v in value.values():
                    yield from colours(v)
            elif isinstance(value, (list, tuple)):
                for v in value:
                    yield from colours(v)

        seen = [c for k, v in mm.STYLES["print"].items() if k != "label"
                for c in colours(v) if isinstance(c, str) and c.startswith("#")]
        self.assertGreater(len(seen), 10)
        for colour in seen:
            r, g, b = hex_rgb(colour)
            self.assertTrue(r == g == b, f"{colour} is not a grey")

    def test_ink_is_pure_black_and_the_land_is_the_white_of_the_page(self):
        style = mm.STYLES["print"]
        self.assertEqual(style["ink"], "#000000")
        self.assertEqual(style["terrain"]["land"], "#ffffff")
        self.assertEqual(style["accent"], "#000000")
        self.assertFalse(style["texture"])
        self.assertFalse(style["vignette"])

    def test_terrains_that_can_touch_differ_in_brightness(self):
        # Breaking it: give the forest nearly the swamp's grey and this names
        # the pair. 0.08 is about 20 steps of 255 - the least a coarse press
        # tint keeps apart.
        fills = self.fills()
        for a, b in self.ADJACENT:
            gap = abs(luma(fills[a]) - luma(fills[b]))
            self.assertGreaterEqual(gap, 0.08, f"{a} and {b} are {gap:.3f} apart")
        # the coast is the one edge that has no ink-free way to show:
        self.assertGreaterEqual(luma(fills["land"]) - luma(fills["sea"]), 0.15)

    def test_a_whole_map_in_it_has_no_colour_in_the_picture(self):
        from PIL import Image, ImageChops

        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 3))
        gm.style = "print"
        with tempfile.TemporaryDirectory() as folder:
            path = mm.render_png(gm, Path(folder) / "p.png", scale=0.5)
            with Image.open(path) as image:
                r, g, b = image.convert("RGB").split()
                self.assertIsNone(ImageChops.difference(r, g).getbbox())
                self.assertIsNone(ImageChops.difference(g, b).getbbox())

    def test_every_style_has_an_accent_and_the_old_ones_are_unchanged(self):
        for name, style in mm.STYLES.items():
            self.assertTrue(style.get("accent", "").startswith("#"), name)
        self.assertEqual(mm.STYLES["parchment"]["accent"], "#8a2f22")
        self.assertEqual(mm.STYLES["treasure"]["accent"], "#8a2f22")
        self.assertEqual(mm.STYLES["ink"]["accent"], mm.STYLES["ink"]["ink"])
        self.assertEqual(mm.STYLES["dark"]["accent"], mm.STYLES["dark"]["ink"])


class PrintSizes(unittest.TestCase):
    def test_the_trim_sizes_a_book_uses(self):
        p = mm.PRINT_PRESETS
        for key, size in (("5x8", (5, 8)), ("5.25x8", (5.25, 8)),
                          ("5.5x8.5", (5.5, 8.5)), ("6x9", (6, 9)),
                          ("letter", (8.5, 11))):
            self.assertEqual((p[key].width_in, p[key].height_in), size, key)
        self.assertAlmostEqual(p["a5"].width_in, 5.827, places=2)
        self.assertAlmostEqual(p["a5"].height_in, 8.268, places=2)
        self.assertAlmostEqual(p["a4"].width_in, 8.268, places=2)
        self.assertAlmostEqual(p["a4"].height_in, 11.693, places=2)

    def test_every_size_has_a_double_page_spread_at_twice_the_width(self):
        singles = [k for k, v in mm.PRINT_PRESETS.items() if not v.spread]
        self.assertEqual(len(singles), 7)
        for key in singles:
            single, spread = mm.PRINT_PRESETS[key], mm.PRINT_PRESETS[f"{key}-spread"]
            self.assertTrue(spread.spread)
            self.assertAlmostEqual(spread.width_in, single.width_in * 2)
            self.assertEqual(spread.height_in, single.height_in)
            self.assertIn("spread", spread.label)

    def test_the_layout_adds_bleed_all_round_and_keeps_the_safe_area_inside_the_trim(self):
        layout = mm.print_layout("6x9", 300)
        self.assertEqual(layout.bleed_px, 38)                  # 0.125 in, rounded up
        self.assertEqual(layout.size, (1800 + 76, 2700 + 76))
        self.assertEqual(layout.trim, (38, 38, 1838, 2738))
        self.assertEqual(layout.safe, (38 + 150, 38 + 150, 1838 - 150, 2738 - 150))
        self.assertIsNone(layout.fold_x)
        spread = mm.print_layout("6x9-spread", 300)
        self.assertEqual(spread.size, (3600 + 76, 2700 + 76))
        self.assertEqual(spread.fold_x, 38 + 1800)

    def test_a_gutter_is_taken_from_the_inside_edge_only(self):
        plain = mm.print_layout("5x8", 100)
        left = mm.print_layout("5x8", 100, gutter_in=0.5, inside="left")
        right = mm.print_layout("5x8", 100, gutter_in=0.5, inside="right")
        self.assertEqual(left.safe[0], plain.safe[0] + 50)
        self.assertEqual(left.safe[2], plain.safe[2])
        self.assertEqual(right.safe[2], plain.safe[2] - 50)
        self.assertEqual(right.safe[0], plain.safe[0])

    def test_nonsense_is_refused_with_a_word_about_what_would_work(self):
        with self.assertRaises(ValueError) as caught:
            mm.print_layout("7x7", 300)
        self.assertIn("6x9", str(caught.exception))
        for kwargs in ({"dpi": 10}, {"dpi": 5000}, {"safe_in": 4.0},
                       {"inside": "middle"}):
            with self.assertRaises(ValueError, msg=str(kwargs)):
                mm.print_layout("6x9", **{"dpi": 300, **kwargs})


class PrintExport(unittest.TestCase):
    DPI = 100                       # the maths is the same at any dpi; this is quick

    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.folder.name)
        cls.count = 0

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def out(self):
        type(self).count += 1
        return self.dir / f"p{self.count}.png"

    def export(self, gm, preset="6x9", **kw):
        kw.setdefault("dpi", self.DPI)
        return mm.export_print(gm, self.out(), preset, **kw)

    def image(self, result):
        from PIL import Image

        image = Image.open(result.path)
        image.load()
        return image

    def ink_box(self, image):
        """Where anything that is not the page's own white sits."""
        from PIL import ImageChops

        grey = image.convert("L")
        return ImageChops.invert(grey).getbbox()

    def test_the_sheet_is_the_trim_plus_bleed_at_the_dpi_and_says_so(self):
        # Breaking it: leave the bleed out of the sheet and the size is short.
        gm = secret_map()
        for key in ("5x8", "6x9", "letter", "a5-spread"):
            layout = mm.print_layout(key, self.DPI)
            result = self.export(gm, key)
            image = self.image(result)
            self.assertEqual(image.size, layout.size, key)
            self.assertEqual((result.width_px, result.height_px), layout.size)
            self.assertEqual(round(image.info["dpi"][0]), self.DPI, key)
            self.assertEqual(round(image.info["dpi"][1]), self.DPI, key)

    def test_the_dpi_is_written_as_asked(self):
        result = self.export(secret_map(), "5x8", dpi=300)
        self.assertEqual(round(self.image(result).info["dpi"][0]), 300)
        self.assertEqual(result.dpi, 300)

    def test_greyscale_is_one_channel_and_colour_is_kept_when_asked(self):
        gm = secret_map()
        self.assertEqual(self.image(self.export(gm)).mode, "L")
        colour = self.image(self.export(gm, greyscale=False))
        self.assertEqual(colour.mode, "RGB")
        r, g, b = colour.split()
        from PIL import ImageChops
        self.assertIsNotNone(ImageChops.difference(r, b).getbbox(),
                             "parchment should still be coloured")

    def test_the_map_sits_inside_the_safe_area_and_the_rest_is_paper(self):
        # Breaking it: fit to the trim instead of the safe area and the map
        # spills into the margin.
        for gm in (secret_map(),
                   mm.GameMap(name="Tall", width=500, height=900)):
            layout = mm.print_layout("6x9", self.DPI)
            image = self.image(self.export(gm))
            box = self.ink_box(image)
            sx0, sy0, sx1, sy1 = layout.safe
            self.assertGreaterEqual(box[0], sx0 - 1)
            self.assertGreaterEqual(box[1], sy0 - 1)
            self.assertLessEqual(box[2], sx1 + 1)
            self.assertLessEqual(box[3], sy1 + 1)
            # as big as the safe area allows, in the shape it was drawn
            self.assertTrue(box[2] - box[0] >= (sx1 - sx0) - 3
                            or box[3] - box[1] >= (sy1 - sy0) - 3, gm.name)
            self.assertAlmostEqual((box[2] - box[0]) / (box[3] - box[1]),
                                   gm.width / gm.height, delta=0.02)
            # centred
            self.assertAlmostEqual((box[0] + box[2]) / 2, (sx0 + sx1) / 2, delta=2)
            self.assertAlmostEqual((box[1] + box[3]) / 2, (sy0 + sy1) / 2, delta=2)

    def test_a_gutter_keeps_the_map_off_the_inside_edge(self):
        gm = secret_map()
        for inside in ("left", "right"):
            layout = mm.print_layout("6x9", self.DPI, gutter_in=0.5, inside=inside)
            box = self.ink_box(self.image(self.export(gm, gutter_in=0.5,
                                                      inside=inside)))
            self.assertGreaterEqual(box[0], layout.safe[0] - 1, inside)
            self.assertLessEqual(box[2], layout.safe[2] + 1, inside)
        plain = self.ink_box(self.image(self.export(gm)))
        gutter = self.ink_box(self.image(self.export(gm, gutter_in=0.5)))
        self.assertLess(gutter[2] - gutter[0], plain[2] - plain[0])

    def test_a_spread_puts_the_map_across_the_spine_and_warns_about_names_on_it(self):
        gm = mm.GameMap(name="Two Pages", width=800, height=550, title_on_map=False,
                        compass=False, layers=[mm.Layer(name="Base")])
        gm.pins.append(mm.Pin(x=400, y=300, kind="town", label="Spine Town"))
        gm.pins.append(mm.Pin(x=120, y=300, kind="town", label="Edge Town"))
        layout = mm.print_layout("6x9-spread", self.DPI)
        result = self.export(gm, "6x9-spread")
        box = self.ink_box(self.image(result))
        self.assertAlmostEqual((box[0] + box[2]) / 2, layout.fold_x, delta=2)
        spine = [n for n in result.notes if "spine" in n]
        self.assertEqual(len(spine), 1)
        self.assertIn("Spine Town", spine[0])
        self.assertNotIn("Edge Town", spine[0])
        quiet = mm.GameMap(name="Quiet", width=800, height=550, title_on_map=False,
                           compass=False, layers=[mm.Layer(name="Base")])
        quiet.pins.append(mm.Pin(x=120, y=300, kind="town", label="Edge Town"))
        self.assertFalse([n for n in self.export(quiet, "6x9-spread").notes
                          if "spine" in n])
        self.assertFalse([n for n in self.export(gm, "6x9").notes if "spine" in n])

    def test_guides_mark_bleed_trim_and_safe_area_and_only_when_asked(self):
        gm = secret_map()
        layout = mm.print_layout("6x9", self.DPI)
        plain = self.image(self.export(gm, greyscale=False))
        proof = self.image(self.export(gm, guides=True))
        self.assertEqual(proof.mode, "RGB")
        self.assertEqual(proof.getpixel((0, 0)), hex_rgb(mm.GUIDE_COLOURS["bleed"]))
        self.assertEqual(proof.getpixel((layout.trim[0], layout.trim[1])),
                         hex_rgb(mm.GUIDE_COLOURS["trim"]))
        self.assertEqual(plain.getpixel((0, 0)), (255, 255, 255))
        self.assertEqual(plain.getpixel((layout.trim[0], layout.trim[1])),
                         (255, 255, 255))
        top_middle = proof.getpixel((layout.trim[0] + 3, layout.safe[1]))
        self.assertIn(top_middle, (hex_rgb(mm.GUIDE_COLOURS["safe"]), (255, 255, 255)))
        spread = self.image(self.export(gm, "6x9-spread", guides=True))
        fold = mm.print_layout("6x9-spread", self.DPI).fold_x
        self.assertEqual(spread.getpixel((fold, 100)), hex_rgb(mm.GUIDE_COLOURS["fold"]))
        self.assertTrue([n for n in self.export(gm, guides=True).notes
                         if "proof" in n])

    def test_the_default_is_the_reader_edition_and_the_authors_secrets_stay_home(self):
        # Breaking it: default `edition` to "author".
        gm = secret_map()
        default = self.image(self.export(gm))
        reader = self.image(self.export(gm, edition="reader"))
        author = self.image(self.export(gm, edition="author"))
        self.assertEqual(default.tobytes(), reader.tobytes())
        self.assertNotEqual(default.tobytes(), author.tobytes())
        result = self.export(gm)
        self.assertEqual(result.alt_text, mm.describe_map(gm, "reader"))
        self.assertNotIn("Vault", result.alt_text)
        self.assertEqual(self.image(result).text.get("Description"), result.alt_text)

    def test_another_style_for_one_export_leaves_the_map_alone(self):
        gm = secret_map()
        plain = self.image(self.export(gm))
        printed = self.image(self.export(gm, style="print"))
        self.assertEqual(gm.style, "parchment")
        self.assertNotEqual(plain.tobytes(), printed.tobytes())
        self.assertFalse([n for n in self.export(gm, style="print").notes
                          if "turned to grey" in n])
        self.assertTrue([n for n in self.export(gm).notes if "turned to grey" in n])

    def test_it_says_how_small_the_names_print(self):
        # 9 px lettering on an 800 px map fitted 5 in wide is 4 pt.
        gm = secret_map()
        result = self.export(gm)
        scale = 500 / 800                 # 6x9 at 100 dpi: safe area 500 px wide
        self.assertAlmostEqual(result.min_text_pt, 9 * scale / self.DPI * 72, delta=0.02)
        self.assertTrue([n for n in result.notes if "smallest names" in n])
        big = mm.GameMap(name="Big", width=800, height=550, title_on_map=False,
                         compass=False, border=False,
                         layers=[mm.Layer(name="Base")])
        big.labels.append(mm.MapLabel(x=400, y=270, text="LARGE", size=60))
        loud = self.export(big)
        self.assertGreater(loud.min_text_pt, 6.0)
        self.assertFalse([n for n in loud.notes if "smallest names" in n])

    def test_no_line_is_thinner_than_the_press_can_hold(self):
        # Breaking it: pass 1 as the minimum stroke to the renderer.
        seen = []
        real = mm._render_image

        def spy(*args, **kw):
            seen.append(args[4] if len(args) > 4 else kw.get("min_stroke"))
            return real(*args, **kw)

        mm._render_image = spy
        try:
            self.export(secret_map(), dpi=300)
            self.export(secret_map(), dpi=300, min_line_pt=0.5)
        finally:
            mm._render_image = real
        self.assertEqual(seen, [2, 3])            # 0.25 pt is 1.04 px; 0.5 pt is 2.08

    def test_a_hairline_is_thickened_when_drawn(self):
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (100, 40), "white")
        prim = ("line", [(5, 20), (95, 20)], "#000000", 0.2, False)
        mm._paint(ImageDraw.Draw(image, "RGBA"), [prim], 1.0, 3)
        rows = [y for y in range(40) if image.getpixel((50, y)) != (255, 255, 255)]
        self.assertEqual(len(rows), 3)
        image = Image.new("RGB", (100, 40), "white")
        mm._paint(ImageDraw.Draw(image, "RGBA"), [prim], 1.0)
        rows = [y for y in range(40) if image.getpixel((50, y)) != (255, 255, 255)]
        self.assertEqual(len(rows), 1)

    def test_a_real_world_exports_quickly_at_300_dpi(self):
        import time

        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 5))
        started = time.perf_counter()
        result = mm.export_print(gm, self.out(), "6x9", style="print")
        self.assertLess(time.perf_counter() - started, 15.0)
        self.assertEqual(self.image(result).size, (1876, 2776))


class EbookExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.folder.name)

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def test_a_colour_png_of_the_width_asked_with_the_description(self):
        from PIL import Image

        gm = secret_map()
        result = mm.export_ebook(gm, self.dir / "e.png", width_px=900)
        self.assertEqual((result.width_px, result.height_px), (900, round(550 * 900 / 800)))
        with Image.open(result.path) as image:
            image.load()
            self.assertEqual(image.size, (result.width_px, result.height_px))
            self.assertEqual(image.mode, "RGB")           # opaque, no alpha
            self.assertEqual(image.text.get("Description"), result.alt_text)
            r, g, b = image.split()
            from PIL import ImageChops
            self.assertIsNotNone(ImageChops.difference(r, b).getbbox(), "kept its colour")
        self.assertEqual(result.alt_text, mm.describe_map(gm, "reader"))
        self.assertNotIn("Vault", result.alt_text)

    def test_the_default_is_the_reader_edition_at_1800_wide(self):
        from PIL import Image

        gm = secret_map()
        result = mm.export_ebook(gm, self.dir / "d.png")
        self.assertEqual(result.width_px, 1800)
        reader = mm.export_ebook(gm, self.dir / "r.png", 1800, edition="reader")
        author = mm.export_ebook(gm, self.dir / "a.png", 1800, edition="author")
        with Image.open(result.path) as d, Image.open(reader.path) as r, \
                Image.open(author.path) as a:
            self.assertEqual(d.tobytes(), r.tobytes())
            self.assertNotEqual(d.tobytes(), a.tobytes())

    def test_a_jpeg_carries_the_description_too(self):
        from PIL import Image

        gm = secret_map()
        result = mm.export_ebook(gm, self.dir / "e.jpg", width_px=700)
        with Image.open(result.path) as image:
            self.assertEqual(image.format, "JPEG")
            self.assertEqual(image.mode, "RGB")
            self.assertEqual(image.info.get("comment"), result.alt_text.encode("utf-8"))

    def test_silly_widths_are_kept_sensible(self):
        gm = secret_map()
        self.assertEqual(mm.export_ebook(gm, self.dir / "s.png", 10).width_px, 300)
        self.assertEqual(mm.export_ebook(gm, self.dir / "b.png", 99999).width_px, 6000)


def keyed_map(**kw):
    """
    A map with a few kinds of thing on it and a legend switched on.

    Base holds land, a forest, a river, a capital and a town. Secret (author
    only) holds a mountain range and a ruin. Hidden (off) holds a desert.
    """
    gm = mm.GameMap(
        name="Key", width=1200, height=825, legend=True, scale_text="20 miles",
        layers=[mm.Layer(name="Base"), mm.Layer(name="Secret", author_only=True),
                mm.Layer(name="Hidden", visible=False)], **kw)
    gm.shapes = [
        mm.Shape(kind="land", layer="Base", points=square(300, 200, 600)),
        mm.Shape(kind="forest", layer="Base", points=square(380, 320, 160)),
        mm.Shape(kind="river", layer="Base", closed=False,
                 points=[(500, 220), (560, 400), (520, 640)]),
        mm.Shape(kind="mountains", layer="Secret", closed=False,
                 points=[(650, 260), (780, 300), (860, 380)]),
        mm.Shape(kind="desert", layer="Hidden", points=square(700, 450, 120)),
    ]
    gm.pins = [
        mm.Pin(x=450, y=500, kind="capital", label="Keep", layer="Base"),
        mm.Pin(x=700, y=600, kind="town", label="Ford", layer="Base"),
        mm.Pin(x=800, y=300, kind="ruin", label="Cairn", layer="Secret"),
    ]
    for i, item in enumerate(gm.shapes + gm.pins):
        item.id = f"key_{i}"
    return gm


def legend_tail(prims):
    """The legend's own primitives: its panel and everything drawn after it."""
    start = next(i for i, p in enumerate(prims) if p[0] == "text" and p[3] == "LEGEND")
    return prims[start - 1:]


class Legend(unittest.TestCase):
    def test_it_is_off_until_asked_for_and_then_it_is_a_key(self):
        self.assertFalse(mm.GameMap().legend)
        gm = keyed_map()
        gm.legend = False
        self.assertNotIn("LEGEND", texts(mm.build_primitives(gm)))
        gm.legend = True
        self.assertIn("LEGEND", texts(mm.build_primitives(gm)))
        self.assertTrue(mm.GameMap.from_json(gm.to_json()).legend)

    def test_it_lists_the_kinds_that_are_drawn_and_no_others(self):
        # Breaking it: list every kind whether it is drawn or not.
        names = [p[3] for p in legend_tail(mm.build_primitives(keyed_map()))
                 if p[0] == "text"]
        for shown in ("Land", "Forest", "River", "Mountains", "Capital city",
                      "Town", "Ruin"):
            self.assertEqual(names.count(shown), 1, shown)
        for absent in ("Desert", "Road", "Wall", "Route", "Hills", "Castle / keep",
                       "Marsh", "Water"):
            self.assertNotIn(absent, names)

    def test_it_lists_only_what_the_edition_shows(self):
        gm = keyed_map()
        author = [p[3] for p in legend_tail(mm.build_primitives(gm)) if p[0] == "text"]
        reader = [p[3] for p in legend_tail(mm.build_primitives(gm, edition="reader"))
                  if p[0] == "text"]
        for secret in ("Mountains", "Ruin"):
            self.assertIn(secret, author)
            self.assertNotIn(secret, reader)
        for name in ("Land", "Forest", "Capital city"):
            self.assertIn(name, reader)
        self.assertNotIn("Desert", author + reader)         # the layer is off

    def test_a_swatch_is_the_colour_the_map_uses(self):
        gm = keyed_map()
        terrain = gm.palette()["terrain"]
        fills = {p[2] for p in legend_tail(mm.build_primitives(gm))
                 if p[0] == "polygon" and p[2]}
        self.assertIn(terrain["land"], fills)
        self.assertIn(terrain["forest"], fills)
        gm.style = "print"
        printed = {p[2] for p in legend_tail(mm.build_primitives(gm))
                   if p[0] == "polygon" and p[2]}
        for colour in printed:
            r, g, b = hex_rgb(colour)
            self.assertTrue(r == g == b, f"{colour} in a print legend is not a grey")

    def test_every_primitive_is_the_kind_all_three_backends_already_draw(self):
        for prim in legend_tail(mm.build_primitives(keyed_map())):
            self.assertIn(prim[0], ("polygon", "line", "ellipse", "text"))
            if prim[0] == "text":
                self.assertIn(len(prim), (10, 11))
            for value in prim[1:]:
                if isinstance(value, str) and value.startswith("#"):
                    self.assertRegex(value, r"^#[0-9a-f]{6}$")

    def test_it_sits_in_a_corner_clear_of_the_title_compass_and_scale(self):
        for gm in (keyed_map(), mapgen.generate(mapgen.preset("Classic fantasy world", 3))):
            gm.legend = True
            layout = mm.legend_layout(gm)
            self.assertIsNotNone(layout)
            x0, y0, x1, y1 = layout.box
            self.assertTrue(0 <= x0 and 0 <= y0 and x1 <= gm.width and y1 <= gm.height)
            for other in mm.furniture_boxes(gm):
                self.assertFalse(mm.boxes_overlap(layout.box, other, 0.0),
                                 f"{gm.name}: the legend is on top of the furniture")

    def test_it_keeps_out_of_the_way_of_the_land(self):
        # Breaking it: ignore what lies under each corner.
        west = mm.GameMap(name="W", width=1200, height=825, legend=True,
                          compass=False, title_on_map=False,
                          layers=[mm.Layer(name="Base")])
        west.shapes = [mm.Shape(id="w", kind="land", layer="Base",
                                points=[(0, 0), (500, 0), (500, 825), (0, 825)])]
        west.pins = [mm.Pin(id="p", x=200, y=400, kind="town", label="A")]
        self.assertIn(mm.legend_layout(west).corner, ("tr", "br"))
        east = mm.GameMap(name="E", width=1200, height=825, legend=True,
                          compass=False, title_on_map=False,
                          layers=[mm.Layer(name="Base")])
        east.shapes = [mm.Shape(id="e", kind="land", layer="Base",
                                points=[(700, 0), (1200, 0), (1200, 825), (700, 825)])]
        east.pins = [mm.Pin(id="p", x=900, y=400, kind="town", label="A")]
        self.assertEqual(mm.legend_layout(east).corner, "tl")
        top = mm.GameMap(name="T", width=1200, height=825, legend=True,
                         compass=False, title_on_map=False,
                         layers=[mm.Layer(name="Base")])
        top.shapes = [mm.Shape(id="t", kind="land", layer="Base",
                               points=[(0, 0), (1200, 0), (1200, 400), (0, 400)])]
        top.pins = [mm.Pin(id="p", x=600, y=200, kind="town", label="A")]
        self.assertIn(mm.legend_layout(top).corner, ("bl", "br"))

    def test_names_are_kept_off_it(self):
        # Breaking it: stop reserving the legend's box when placing names.
        gm = mm.GameMap(name="Key", width=1200, height=825, legend=True,
                        compass=False, title_on_map=False,
                        layers=[mm.Layer(name="Base")])
        gm.shapes = [mm.Shape(id="l", kind="land", layer="Base",
                              points=square(600, 300, 300))]
        gm.pins = [mm.Pin(id="near", x=0, y=0, kind="town", label="Nearby",
                          label_side="w")]
        box = mm.legend_layout(gm).box
        gm.pins[0].x, gm.pins[0].y = box[2] + 24, box[1] + 30       # just outside it
        self.assertEqual(mm.legend_layout(gm).box, box)
        gm.legend = False
        self.assertEqual(mm.layout_pin_labels(gm)["near"], ("w", True))
        gm.legend = True
        side, shown = mm.layout_pin_labels(gm)["near"]
        self.assertNotEqual((side, shown), ("w", True),
                            "the name was placed on top of the legend")

    def test_the_remembered_name_placement_notices_a_legend_arriving(self):
        # Breaking it: leave the legend's box out of the placement memo's key -
        # switching the legend on then serves the old, overlapping placement.
        gm = mm.GameMap(name="Key", width=1200, height=825, legend=False,
                        compass=False, title_on_map=False,
                        layers=[mm.Layer(name="Base")])
        gm.shapes = [mm.Shape(id="l", kind="land", layer="Base",
                              points=square(600, 300, 300))]
        gm.pins = [mm.Pin(id="near", x=0, y=0, kind="town", label="Nearby",
                          label_side="w")]
        gm.legend = True
        box = mm.legend_layout(gm).box
        gm.legend = False
        gm.pins[0].x, gm.pins[0].y = box[2] + 24, box[1] + 30

        def nearby(prims):
            return [(p[1], p[2], p[6]) for p in prims
                    if p[0] == "text" and p[3] == "Nearby"]

        before = nearby(mm.build_primitives(gm))
        self.assertEqual(before[0][2], "e")                # west of the pin
        gm.legend = True
        self.assertNotEqual(nearby(mm.build_primitives(gm)), before)

    def test_a_busy_world_gets_a_compact_legend(self):
        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 3))
        gm.legend = True
        layout = mm.legend_layout(gm)
        self.assertLessEqual(len(layout.entries), mm.MAX_LEGEND_ENTRIES)
        self.assertLessEqual(layout.rows, mm.LEGEND_ROWS)
        x0, y0, x1, y1 = layout.box
        self.assertLess((x1 - x0) * (y1 - y0), 0.3 * gm.width * gm.height)
        crowded = keyed_map()
        for i, kind in enumerate(list(mm.PIN_KINDS) + list(mm.EXTRA_PIN_KINDS)):
            crowded.pins.append(mm.Pin(id=f"c{i}", x=310 + i * 20, y=700, kind=kind,
                                       layer="Base"))
        entries = mm.legend_entries(crowded)
        self.assertEqual(len(entries), mm.MAX_LEGEND_ENTRIES)
        self.assertEqual([e[0] for e in entries[:3]], ["terrain"] * 3)   # ground first

    def test_an_empty_map_has_nothing_to_list(self):
        gm = mm.GameMap(name="Blank", legend=True)
        self.assertIsNone(mm.legend_layout(gm))
        self.assertNotIn("LEGEND", texts(mm.build_primitives(gm)))

    def test_switching_it_on_and_off_rebuilds_and_a_repeat_look_does_not(self):
        # Breaking it: leave `legend` out of the display list's key.
        gm = keyed_map()
        gm.legend = False
        without = mm.build_primitives(gm)
        gm.legend = True
        with_it = mm.build_primitives(gm)
        self.assertNotEqual(without, with_it)
        cached = gm._prim_cache
        mm.build_primitives(gm)
        self.assertIs(gm._prim_cache, cached, "a second look rebuilt everything")
        gm.legend = False
        self.assertEqual(mm.build_primitives(gm), without)

    def test_the_svg_and_the_picture_carry_it_too(self):
        from PIL import Image

        gm = keyed_map()
        with tempfile.TemporaryDirectory() as folder:
            svg = mm.render_svg(gm, Path(folder) / "k.svg").read_text(encoding="utf-8")
            for name in ("LEGEND", "Forest", "Capital city"):
                self.assertIn(name, svg)
            on = mm.render_png(gm, Path(folder) / "on.png", scale=0.5)
            gm.legend = False
            off = mm.render_png(gm, Path(folder) / "off.png", scale=0.5)
            with Image.open(on) as a, Image.open(off) as b:
                self.assertNotEqual(a.convert("RGB").tobytes(), b.convert("RGB").tobytes())

    def test_working_out_the_legend_is_quick(self):
        import time

        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 5))
        gm.legend = True
        started = time.perf_counter()
        for _ in range(20):
            mm.legend_layout(gm)
        self.assertLess((time.perf_counter() - started) / 20, 0.05)

    def test_a_legend_does_not_change_a_map_that_has_none(self):
        gm = mapgen.generate(mapgen.preset("Classic fantasy world", 3))
        before = mm.build_primitives(gm)
        gm.legend = True
        gm.legend = False
        self.assertEqual(mm.build_primitives(gm), before)


def nested_maps():
    """A world, a village on it, and an inn's floor plan inside the village."""
    world = mm.GameMap(id="map_world", name="The Known World", seed=7,
                       layers=[mm.Layer(name="Base")])
    village = mm.GameMap(id="map_village", name="Harrowgate", seed=11, kind="city",
                         layers=[mm.Layer(name="Base")])
    inn = mm.GameMap(id="map_inn", name="The Gilded Stag, ground floor", seed=12,
                     kind="building", layers=[mm.Layer(name="Base")])
    world.pins.append(mm.Pin(id="pin_h", x=100, y=100, label="Harrowgate Town",
                             kind="town"))
    village.pins.append(mm.Pin(id="pin_i", x=50, y=50, label="Gilded Stag", kind="inn"))
    mm.link_child(world, world.pins[0], village)
    mm.link_child(village, village.pins[0], inn)
    return world, village, inn


class MapsInsideMaps(unittest.TestCase):
    def test_a_child_seed_is_a_crc_of_the_parent_seed_and_the_pin(self):
        # Breaking it: seed with hash(), which Python salts differently in every
        # process, and the village looks different each time the book is opened.
        self.assertEqual(mm.child_seed(7, "pin_abc"), 80623281)
        self.assertEqual(mm.child_seed(2024, "pin_0001"), 3059890015)
        self.assertEqual(mm.child_seed(7, "pin_abc"),
                         zlib.crc32(b"7:pin_abc") & 0xFFFFFFFF)
        seeds = {mm.child_seed(s, p) for s in (1, 2, 3) for p in ("a", "b", "c")}
        self.assertEqual(len(seeds), 9)                    # a different place, a different map
        self.assertTrue(all(isinstance(s, int) and 0 <= s < 2 ** 32 for s in seeds))
        self.assertIs(mapstory.child_seed, mm.child_seed)

    def test_it_is_the_same_in_every_process(self):
        import os
        import subprocess
        import sys

        from tests import REPO, SANDBOX

        code = ("import sys; sys.path.insert(0, sys.argv[1]); "
                "from novelforge import mapmaker as mm; "
                "print(mm.child_seed(7, 'pin_abc'), mm.child_seed(99, 'x'))")
        answers = set()
        for salt in ("1", "2", "random"):
            env = dict(os.environ, PYTHONHASHSEED=salt,
                       NOVELFORGE_SETTINGS=str(SANDBOX / "s.json"),
                       NOVELFORGE_PROJECTS=str(SANDBOX / "p"))
            done = subprocess.run([sys.executable, "-c", code, str(REPO)], env=env,
                                  capture_output=True, text=True, timeout=60)
            self.assertEqual(done.returncode, 0, done.stderr)
            answers.add(done.stdout.strip())
        self.assertEqual(len(answers), 1, answers)
        self.assertEqual(answers.pop().split()[0], "80623281")

    def test_moving_or_renaming_a_pin_does_not_change_the_map_inside_it(self):
        pin = mm.Pin(id="pin_x", x=10, y=10, label="Old name")
        first = mm.child_seed(7, pin.id)
        pin.x, pin.y, pin.label, pin.kind = 400, 300, "New name", "city"
        self.assertEqual(mm.child_seed(7, pin.id), first)

    def test_the_links_are_saved_and_come_back_and_old_files_do_not_mind(self):
        world, village, inn = nested_maps()
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            for gm in (world, village, inn):
                mm.save_map(folder, gm)
            again = {m.id: m for m in (mm.load_map(p) for _n, p in mm.list_maps(folder))}
        self.assertEqual(again["map_world"].pins[0].child_map_id, "map_village")
        self.assertEqual(again["map_village"].parent,
                         {"map_id": "map_world", "pin_id": "pin_h"})
        self.assertEqual(again["map_inn"].parent,
                         {"map_id": "map_village", "pin_id": "pin_i"})
        self.assertEqual(again["map_world"].parent, {})
        old_pin = mm.GameMap.from_json({"pins": [{"id": "p", "x": 1, "y": 2}]}).pins[0]
        self.assertEqual(old_pin.child_map_id, "")
        self.assertEqual(mm.GameMap.from_json({"parent": "garbage"}).parent, {})
        self.assertEqual(mm.GameMap().parent, {})

    def test_two_new_maps_never_share_a_parent_dict(self):
        a, b = mm.GameMap(), mm.GameMap()
        a.parent["map_id"] = "x"
        self.assertEqual(b.parent, {})

    def test_linking_records_both_sides_and_unlinking_undoes_it(self):
        world, village, _inn = nested_maps()
        pin = world.pins[0]
        self.assertEqual(pin.child_map_id, village.id)
        self.assertEqual(village.parent, {"map_id": world.id, "pin_id": pin.id})
        mm.unlink_child(world, pin, village)
        self.assertEqual((pin.child_map_id, village.parent), ("", {}))
        with self.assertRaises(ValueError):
            mm.link_child(world, mm.Pin(id="stranger"), village)   # not on that map
        with self.assertRaises(ValueError):
            mm.link_child(world, pin, world)                       # inside itself

    def test_the_breadcrumb_runs_from_the_top_down_to_this_map(self):
        # Breaking it: build the chain bottom-up.
        world, village, inn = nested_maps()
        maps = [inn, world, village]                       # any order
        crumbs = mapstory.breadcrumb(maps, inn)
        self.assertEqual([c.name for c in crumbs],
                         ["The Known World", "Harrowgate",
                          "The Gilded Stag, ground floor"])
        self.assertEqual([c.map_id for c in crumbs],
                         ["map_world", "map_village", "map_inn"])
        self.assertEqual([c.pin_id for c in crumbs], ["", "pin_h", "pin_i"])
        self.assertEqual([c.pin_label for c in crumbs],
                         ["", "Harrowgate Town", "Gilded Stag"])
        self.assertEqual(mapstory.breadcrumb_text(maps, inn),
                         "The Known World > Harrowgate > The Gilded Stag, ground floor")
        self.assertEqual(mapstory.breadcrumb_text(maps, village, " / "),
                         "The Known World / Harrowgate")
        self.assertEqual([c.name for c in mapstory.breadcrumb(maps, world)],
                         ["The Known World"])

    def test_a_missing_parent_or_a_loop_still_gives_a_breadcrumb(self):
        world, village, inn = nested_maps()
        # the world map was deleted: the chain starts at the village
        self.assertEqual([c.name for c in mapstory.breadcrumb([village, inn], inn)],
                         ["Harrowgate", "The Gilded Stag, ground floor"])
        # a map that is not in the list at all (never saved) is its own top
        stranger = mm.GameMap(id="map_new", name="Unsaved")
        self.assertEqual([c.name for c in mapstory.breadcrumb([world], stranger)],
                         ["Unsaved"])
        # a loop, however it got there, is cut
        world.parent = {"map_id": "map_inn", "pin_id": "x"}
        names = [c.name for c in mapstory.breadcrumb([world, village, inn], inn)]
        self.assertEqual(sorted(names), sorted(["The Known World", "Harrowgate",
                                                "The Gilded Stag, ground floor"]))
        self.assertEqual(len(names), 3)


FILLED = ("building", "ward", "plaza", "room", "floor", "stairs", "cave")
LINES = ("street", "lane", "partition", "door", "window")
ORIGINAL_ORDER = ["land", "water", "forest", "mountains", "hills", "desert",
                  "swamp", "ice", "region", "river", "road", "wall", "route"]


def circle(cx, cy, r, n=48):
    return [(cx + r * math.cos(2 * math.pi * i / n),
             cy + r * math.sin(2 * math.pi * i / n)) for i in range(n)]


def kinds_map(style="parchment"):
    """
    One shape of every generator-only kind, each in a cell of its own, and one
    pin of every kind along the bottom. Lines are 10 wide so a pixel taken from
    the middle of one is the line's colour and nothing else.
    """
    gm = mm.GameMap(name="Every kind", width=1200, height=1000, style=style,
                    compass=False, border=False, title_on_map=False,
                    layers=[mm.Layer(name="Base")])
    for n, kind in enumerate(mm.GENERATOR_KINDS):
        cx, cy = 190 + (n % 4) * 280, 150 + (n // 4) * 200
        if kind in FILLED:
            points, extra = square(cx - 50, cy - 50, 100), {"width": 2.0}
            if kind == "cave":                       # organic: not a box
                points = [(cx - 60, cy), (cx - 20, cy - 55), (cx + 30, cy - 50),
                          (cx + 60, cy + 5), (cx + 15, cy + 55), (cx - 40, cy + 45)]
        elif kind == "orbit":
            points, extra = circle(cx, cy, 80), {"width": 6.0}
        elif kind in ("door", "window"):
            points, extra = [(cx - 40, cy), (cx + 40, cy)], {"width": 10.0}
        else:
            points, extra = [(cx - 90, cy), (cx, cy + 6), (cx + 90, cy)], {"width": 10.0}
        gm.shapes.append(mm.Shape(id=f"kind_{kind}", kind=kind, layer="Base",
                                  closed=kind in FILLED or kind == "orbit",
                                  points=points, **extra))
    for i, kind in enumerate(mm.ALL_PIN_KINDS):
        gm.pins.append(mm.Pin(id=f"pin_{kind}", x=60 + i * 46, y=960, kind=kind,
                              layer="Base"))
    return gm


def shape_prims(prims, shape):
    """The primitives a shape drew, found by where its points are."""
    xs = [p[0] for p in shape.points]
    ys = [p[1] for p in shape.points]
    box = (min(xs) - 40, min(ys) - 40, max(xs) + 40, max(ys) + 40)
    found = []
    for prim in prims:
        if prim[0] in ("polygon", "line") and prim[1]:
            inside = all(box[0] <= x <= box[2] and box[1] <= y <= box[3]
                         for x, y in prim[1])
            if inside:
                found.append(prim)
    return found


class NewKindRegistry(unittest.TestCase):
    """The kinds only the generators make: towns, floor plans, dungeons, star maps."""

    def test_the_editors_own_lists_are_exactly_what_they_were(self):
        # The editor lists TERRAIN_ORDER and PIN_KINDS to the writer, who should
        # not be offered a "door" or a "planet"; the site also states both counts.
        self.assertEqual(mm.TERRAIN_ORDER, ORIGINAL_ORDER)
        self.assertEqual(len(mm.PIN_KINDS), 20)
        self.assertEqual(set(mm.TERRAIN) - set(mm.TERRAIN_ORDER), set(mm.GENERATOR_KINDS))
        self.assertEqual(len(mm.GENERATOR_KINDS), 13)
        self.assertFalse(set(mm.EXTRA_PIN_KINDS) & set(mm.PIN_KINDS))
        self.assertEqual(mm.ALL_PIN_KINDS, {**mm.PIN_KINDS, **mm.EXTRA_PIN_KINDS})
        self.assertEqual(set(mm.EXTRA_PIN_KINDS), {"star", "planet", "station", "gate"})

    def test_every_kind_is_registered_with_no_decoration(self):
        for kind in mm.GENERATOR_KINDS:
            spec = mm.TERRAIN[kind]
            self.assertTrue(spec["label"], kind)
            self.assertIsNone(spec["decor"], f"{kind} must be a plain fill or line")
            self.assertIn(kind, mm.PAINT_ORDER, kind)
            self.assertIsInstance(spec["closed"], bool)
            self.assertFalse(spec["halo"], kind)
            self.assertEqual(bool(spec["fill"]), kind in FILLED, kind)
        self.assertEqual({k for k in mm.PAINT_ORDER if k not in ORIGINAL_ORDER},
                         set(mm.GENERATOR_KINDS))

    def test_the_legend_has_a_name_and_a_place_in_the_list_for_every_kind(self):
        for kind in mm.GENERATOR_KINDS:
            self.assertIn(kind, mm.LEGEND_NAMES, kind)
            self.assertIn(kind, mm.LEGEND_ORDER, kind)

    def test_things_are_painted_in_a_sensible_order(self):
        order = mm.PAINT_ORDER
        for lower, higher in (("land", "floor"), ("floor", "room"),
                              ("room", "partition"), ("partition", "door"),
                              ("door", "window"), ("land", "ward"), ("ward", "plaza"),
                              ("plaza", "street"), ("street", "building"),
                              ("road", "building"), ("building", "wall"),
                              ("floor", "stairs"), ("land", "cave"),
                              ("forest", "room")):
            self.assertLess(order[lower], order[higher], f"{lower} under {higher}")
        for kind in ORIGINAL_ORDER:                    # nothing old has moved
            self.assertEqual(order[kind], dict(
                water=1, land=1, ice=2, desert=3, swamp=4, forest=5, hills=6,
                mountains=7, region=8, river=9, road=10, wall=11, route=12)[kind])

    def test_every_style_has_a_colour_for_every_kind_and_an_accent(self):
        for name, style in mm.STYLES.items():
            for kind in mm.GENERATOR_KINDS:
                if kind == "orbit" or kind in LINES or kind in FILLED:
                    self.assertRegex(style["terrain"][kind], r"^#[0-9a-f]{6}$",
                                     f"{name}/{kind}")
            self.assertRegex(style["accent"], r"^#[0-9a-f]{6}$", name)
        for name in ("parchment", "ink", "dark", "treasure"):
            terrain = mm.STYLES[name]["terrain"]
            for kind in ("building", "ward", "room", "floor", "stairs", "cave"):
                self.assertNotEqual(terrain[kind], terrain["land"], f"{name}/{kind}")
            self.assertNotEqual(terrain["building"], terrain["room"], name)
        self.assertEqual(mm.STYLES["parchment"]["terrain"]["door"],
                         mm.STYLES["parchment"]["accent"])        # doors take the accent

    def test_in_the_print_style_neighbouring_kinds_are_a_step_apart_and_all_grey(self):
        terrain = mm.STYLES["print"]["terrain"]
        fills = {**terrain, "sea": mm.STYLES["print"]["paper"]}
        touching = (("land", "ward"), ("ward", "building"), ("ward", "plaza"),
                    ("plaza", "building"), ("land", "building"), ("sea", "building"),
                    ("sea", "floor"), ("floor", "room"), ("floor", "stairs"),
                    ("room", "stairs"), ("sea", "room"), ("sea", "stairs"),
                    ("sea", "cave"), ("cave", "stairs"), ("floor", "cave"),
                    ("sea", "ward"))
        for a, b in touching:
            gap = abs(luma(fills[a]) - luma(fills[b]))
            self.assertGreaterEqual(gap, 0.08, f"{a} and {b} are {gap:.3f} apart")
        for kind in mm.GENERATOR_KINDS:
            r, g, b = hex_rgb(terrain[kind])
            self.assertTrue(r == g == b, f"print {kind} {terrain[kind]} is not a grey")

    def test_the_new_pins_are_drawn_and_are_not_the_default_diamond(self):
        landmark = mm.pin_primitives(mm.Pin(x=50, y=50, kind="landmark", size=8),
                                     "#000000", "#ffffff", "#ff0000")
        seen = []
        for kind in mm.ALL_PIN_KINDS:
            prims = mm.pin_primitives(mm.Pin(x=50, y=50, kind=kind, size=8),
                                      "#000000", "#ffffff", "#ff0000")
            self.assertTrue(prims, kind)
            coords = [v for p in prims if p[0] in ("polygon", "line")
                      for point in p[1] for v in point]
            coords += [v for p in prims if p[0] == "ellipse" for v in p[1:5]]
            reach = max(abs(v - 50) for v in coords)
            self.assertLess(reach, 8 * 3.0, f"{kind} sprawls")
            if kind in mm.EXTRA_PIN_KINDS:
                self.assertNotEqual(prims, landmark, f"{kind} fell back to the diamond")
                seen.append(prims)
        for i, a in enumerate(seen):
            for b in seen[i + 1:]:
                self.assertNotEqual(a, b, "two of the new pins look the same")
        for kind in mm.EXTRA_PIN_KINDS:
            self.assertIn(kind, mm.LABEL_PRIORITY)
            self.assertEqual(mm.pin_kind_label(kind), mm.EXTRA_PIN_KINDS[kind])
        self.assertEqual(mm.pin_kind_label("castle"), "Castle / keep")
        self.assertEqual(mm.pin_kind_label("nonsense"), "nonsense")

    def test_each_kind_of_map_is_offered_only_the_pins_that_belong_on_it(self):
        for kind in list(mm.MAP_KINDS) + ["sector", "system", "castle", "journey"]:
            offered = mm.pin_kinds_for(kind)
            self.assertTrue(offered, kind)
            self.assertEqual(len(offered), len(set(offered)), kind)
            for pin in offered:
                self.assertIn(pin, mm.ALL_PIN_KINDS, f"{kind}: {pin}")
        for kind in ("world", "continent", "region", "unheard-of"):
            self.assertEqual(mm.pin_kinds_for(kind), list(mm.PIN_KINDS), kind)
        for kind in ("world", "continent", "region"):        # no stars over a kingdom
            self.assertFalse(set(mm.pin_kinds_for(kind)) & set(mm.EXTRA_PIN_KINDS))
        for kind in ("sector", "system"):
            self.assertEqual(mm.pin_kinds_for(kind)[:2], ["star", "planet"])
            self.assertNotIn("village", mm.pin_kinds_for(kind))
            self.assertNotIn("capital", mm.pin_kinds_for(kind))
        for kind in ("building", "dungeon", "battle", "treasure"):
            self.assertNotIn("capital", mm.pin_kinds_for(kind))
            self.assertNotIn("village", mm.pin_kinds_for(kind))
            self.assertNotIn("star", mm.pin_kinds_for(kind))
        self.assertIn("gate", mm.pin_kinds_for("city"))
        self.assertIn("dungeon", mm.pin_kinds_for("dungeon"))
        mm.pin_kinds_for("world").append("junk")               # a copy, not the table
        self.assertNotIn("junk", mm.pin_kinds_for("world"))


class NewKindsRender(unittest.TestCase):
    """A map holding every new kind draws in the editor's canvas, the PNG and the SVG."""

    STYLE_NAMES = ("parchment", "ink", "dark", "treasure", "print")

    def test_every_shape_produces_primitives_in_the_styles_own_colours(self):
        # Breaking it: drop a kind from the style palette, or draw a generator
        # line in the ink instead of the palette's colour.
        for style in self.STYLE_NAMES:
            gm = kinds_map(style)
            terrain = gm.palette()["terrain"]
            prims = mm.build_primitives(gm)
            for shape in gm.shapes:
                mine = shape_prims(prims, shape)
                self.assertTrue(mine, f"{style}/{shape.kind} drew nothing")
                if shape.kind in FILLED:
                    fills = [p for p in mine if p[0] == "polygon" and p[2]]
                    self.assertTrue(fills, f"{style}/{shape.kind} has no filled polygon")
                    self.assertEqual(fills[0][2], terrain[shape.kind],
                                     f"{style}/{shape.kind}")
                    self.assertTrue(any(p[0] == "line" for p in mine),
                                    f"{style}/{shape.kind} has no outline")
                else:
                    lines = [p for p in mine if p[0] == "line"]
                    self.assertTrue(lines, f"{style}/{shape.kind} has no line")
                    self.assertEqual(lines[0][2], terrain[shape.kind],
                                     f"{style}/{shape.kind}")
                    self.assertEqual(lines[0][3], shape.width)
                    self.assertEqual(bool(lines[0][4]), shape.kind == "orbit")

    def test_walls_rooms_and_buildings_keep_their_corners(self):
        # Breaking it: smooth every shape. Chaikin corner-cutting turns a
        # rectangle into an octagon with rounded ends.
        gm = kinds_map()
        prims = mm.build_primitives(gm)
        for kind in ("building", "room", "floor", "stairs", "ward", "plaza"):
            shape = gm.shape(f"kind_{kind}")
            polygon = next(p for p in shape_prims(prims, shape)
                           if p[0] == "polygon" and p[2])
            self.assertEqual(polygon[1], shape.points + [shape.points[0]], kind)
        for kind in ("partition", "street"):
            shape = gm.shape(f"kind_{kind}")
            line = next(p for p in shape_prims(prims, shape) if p[0] == "line")
            self.assertEqual(line[1], shape.points, kind)
        cave = gm.shape("kind_cave")
        rounded = next(p for p in shape_prims(prims, cave) if p[0] == "polygon" and p[2])
        self.assertGreater(len(rounded[1]), len(cave.points) + 1)     # organic: smoothed

    def test_an_orbit_is_a_dashed_ring_with_nothing_to_fill(self):
        gm = kinds_map()
        orbit = gm.shape("kind_orbit")
        mine = shape_prims(mm.build_primitives(gm), orbit)
        self.assertFalse([p for p in mine if p[0] == "polygon" and p[2]])
        ring = next(p for p in mine if p[0] == "line")
        self.assertTrue(ring[4])
        self.assertEqual(ring[1][0], ring[1][-1])                     # closed

    def test_an_ordinary_shape_left_open_keeps_its_inked_outline(self):
        # The lines above are coloured from the palette; a plain land, water or
        # forest shape left open must not pick that up.
        gm = mm.GameMap(name="Open", width=1200, height=300, compass=False,
                        border=False, title_on_map=False,
                        layers=[mm.Layer(name="Base")])
        for n, kind in enumerate(("land", "water", "forest")):
            x = n * 400
            gm.shapes.append(mm.Shape(id=kind, kind=kind, closed=False,
                                      points=[(x + 20, 20), (x + 200, 40),
                                              (x + 300, 200)]))
        palette = gm.palette()
        ink = palette["ink"]
        by_kind = {}
        for shape in gm.shapes:
            lines = [p for p in shape_prims(mm.build_primitives(gm), shape)
                     if p[0] == "line"]
            by_kind[shape.kind] = max(lines, key=lambda p: len(p[1]))[2]
        self.assertEqual(by_kind["land"], ink)
        self.assertEqual(by_kind["forest"], ink)
        self.assertEqual(by_kind["water"], mm._mix(ink, palette["terrain"]["water"], 0.5))

    def test_the_svg_is_well_formed_and_holds_every_kind(self):
        import xml.etree.ElementTree as ET

        for style in ("parchment", "print"):
            gm = kinds_map(style)
            terrain = gm.palette()["terrain"]
            with tempfile.TemporaryDirectory() as folder:
                path = mm.render_svg(gm, Path(folder) / "k.svg")
                root = ET.parse(path).getroot()
                text = path.read_text(encoding="utf-8")
            ns = "{http://www.w3.org/2000/svg}"
            polygons = root.findall(f".//{ns}polygon")
            lines = root.findall(f".//{ns}polyline")
            self.assertGreater(len(polygons), len(FILLED))
            self.assertGreater(len(lines), len(LINES))
            for kind in FILLED:
                self.assertIn(f'fill="{terrain[kind]}"', text, f"{style}/{kind}")
            for kind in LINES + ("orbit",):
                self.assertIn(f'stroke="{terrain[kind]}"', text, f"{style}/{kind}")
            self.assertIn("stroke-dasharray", text)            # the orbit
            shape = gm.shape("kind_building")
            first = " ".join(f"{x:.2f},{y:.2f}" for x, y in shape.points)
            self.assertIn(first, text)                          # sharp corners, as given

    def test_the_picture_has_every_kind_where_it_was_drawn(self):
        from PIL import Image

        def close(a, b):
            return all(abs(x - y) <= 3 for x, y in zip(a, b))

        for style in ("parchment", "dark", "print"):
            gm = kinds_map(style)
            terrain = gm.palette()["terrain"]
            with tempfile.TemporaryDirectory() as folder:
                png = mm.render_png(gm, Path(folder) / "k.png")
                with Image.open(png) as image:
                    image = image.convert("RGB")
                    for shape in gm.shapes:
                        want = mm._rgb(terrain[shape.kind])
                        if shape.kind in FILLED:
                            cx = sum(p[0] for p in shape.points) / len(shape.points)
                            cy = sum(p[1] for p in shape.points) / len(shape.points)
                            got = image.getpixel((int(cx), int(cy)))
                            self.assertTrue(close(got, want),
                                            f"{style}/{shape.kind}: {got} != {want}")
                        elif shape.kind == "orbit":
                            hits = sum(1 for x, y in shape.points
                                       if close(image.getpixel((int(x), int(y))), want))
                            self.assertGreater(hits, len(shape.points) * 0.3,
                                               f"{style}/orbit is not on the page")
                        else:
                            (x0, y0), (x1, y1) = shape.points[:2]     # first stretch
                            got = image.getpixel((int((x0 + x1) / 2),
                                                  int((y0 + y1) / 2)))
                            self.assertTrue(close(got, want),
                                            f"{style}/{shape.kind}: {got} != {want}")

    def test_the_pins_and_the_new_shapes_are_saved_and_come_back_drawn_the_same(self):
        gm = kinds_map()
        again = mm.GameMap.from_json(gm.to_json())
        self.assertEqual(mm.build_primitives(again), mm.build_primitives(gm))
        self.assertEqual([p.kind for p in again.pins], list(mm.ALL_PIN_KINDS))

    def test_the_legend_and_the_alt_text_know_the_new_kinds(self):
        gm = kinds_map()
        gm.legend = True
        names = [e[2] for e in mm.legend_entries(gm)]
        for shown in ("Building", "District", "Street", "Room", "Interior wall", "Stairs"):
            self.assertIn(shown, names)
        prims = mm.build_primitives(gm)                      # and it draws without a hitch
        self.assertIn("LEGEND", texts(prims))
        alt = mm.describe_map(gm)
        self.assertIn("buildings", alt)
        self.assertIn("streets", alt)

    def test_a_city_of_a_thousand_buildings_is_still_quick(self):
        import time

        gm = mm.GameMap(name="Bigtown", width=2400, height=1650, compass=False,
                        border=False, title_on_map=False,
                        layers=[mm.Layer(name="Base")])
        for i in range(1000):
            x, y = 80 + (i % 40) * 56, 80 + (i // 40) * 60
            gm.shapes.append(mm.Shape(id=f"b{i}", kind="building",
                                      points=square(x, y, 40)))
        started = time.perf_counter()
        prims = mm.build_primitives(gm)
        cold = time.perf_counter() - started
        self.assertGreater(len(prims), 1000)
        self.assertLess(cold, 1.5)
        started = time.perf_counter()
        mm.build_primitives(gm)
        self.assertLess(time.perf_counter() - started, 0.25)     # remembered


class NewKindsOnTheCanvas(unittest.TestCase):
    """
    The editor's own drawing code, run over a hidden canvas: every primitive a map
    of new kinds produces must put something on it. (The editor swallows a Tk
    error and draws nothing, so a bad colour would otherwise go unseen.)
    """

    def test_every_primitive_reaches_the_canvas(self):
        import tkinter as tk
        from types import SimpleNamespace

        try:
            root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"no display: {error}")
        try:
            root.withdraw()
            from novelforge.ui import mapeditor

            for style in ("parchment", "dark", "print"):
                canvas = tk.Canvas(root, width=1200, height=1000)
                stub = SimpleNamespace(canvas=canvas, zoom=1.0,
                                       to_screen=lambda x, y: (x, y))
                prims = mm.build_primitives(kinds_map(style))
                self.assertGreater(len(prims), 60)
                for prim in prims:
                    before = len(canvas.find_all())
                    mapeditor.MapEditor._draw_primitive(stub, prim)
                    drew = len(canvas.find_all()) - before
                    if prim[0] == "polygon" and len(prim[1]) >= 3:
                        self.assertGreaterEqual(drew, 1, f"{style}: {prim[:4]}")
                    elif prim[0] == "line" and len(prim[1]) >= 2 and prim[2]:
                        self.assertGreaterEqual(drew, 1, f"{style}: {prim[:3]}")
                    elif prim[0] == "ellipse":
                        self.assertGreaterEqual(drew, 1 if abs(prim[3] - prim[1]) >= 1
                                                and abs(prim[4] - prim[2]) >= 1 else 0,
                                                f"{style}: {prim[:5]}")
                canvas.destroy()
        finally:
            root.destroy()


if __name__ == "__main__":
    unittest.main()
